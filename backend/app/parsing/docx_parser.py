"""DOCX 解析器：走 Heading 样式、表格与封面元数据。"""

import re
from datetime import date, datetime
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsing.contract import (
    ClauseNode,
    DocumentMeta,
    ParsedDocument,
    ParseError,
    ParseWarning,
    SourceBlock,
)
from app.parsing.docx_numbering import HeadingNumbering
from app.parsing.headings import acts_as_headings, demote_fragment_headings
from app.parsing.numbering import _noise
from app.parsing.unstructured import append_text, structure_warnings
from app.parsing.vocab import parse_label

_HEADING = re.compile(r"^Heading (\d+)$")
_META = re.compile(
    r"^\s*(Version|Owner|Approver|Approved Date|Effective Date)\s*:\s*(.+?)\s*$"
)
_DOC_TYPE = re.compile(r"\bIT\s+(Policy|Standard|Procedure|Guideline)\b", re.IGNORECASE)
_DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d %B %Y", "%B %d, %Y")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_TEXT = f"{_W}t"
_DELETED = f"{_W}del"


def _element_text(element) -> str:
    """收集段落/单元格里的 w:t，跳过 w:del 下的内容。

    不能用 python-docx 的 Paragraph.text：它只拼接 w:p 的**直接** w:r 子节点，
    嵌在 w:ins（修订插入）里的 run 会被静默丢弃。实测一份带修订痕迹的真实
    文档丢了 96% 的插入片段（约 4700 字符），句子中间凭空少词。

    下游没有任何闸门能察觉这件事：引文确实逐字存在于（已残缺的）条款正文里，
    模型却会自行补上一个看似合理的主语，读起来像有据可依。

    w:del 的正文放在 w:delText，本来就取不到；这里显式跳过是防个别写法把
    w:t 塞进 w:del——修订删除的文字不该出现在正文中。
    """
    parts: list[str] = []

    def walk(node) -> None:
        for child in node:
            if child.tag == _DELETED:
                continue
            if child.tag == _TEXT:
                parts.append(child.text or "")
            else:
                walk(child)

    walk(element)
    return "".join(parts)


def _paragraph_text(paragraph: Paragraph) -> str:
    return _element_text(paragraph._element)


def _mark_deleted(paragraph: Paragraph) -> bool:
    """段落标记在修订中被删除（`w:pPr/w:rPr/w:del`）。

    Word 按最终状态显示时，这一段并进下一段，本身不存在，也不占标题编号。
    不认出来的话，删掉的空标题照样推进编号计数器，导致后续编号错位。
    """
    ppr = paragraph._p.pPr
    rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    return rpr is not None and rpr.find(qn("w:del")) is not None


def _heading_level(paragraph: Paragraph) -> int | None:
    """标题样式给的层级；没有时看段落自己设的大纲级别。

    Word 认标题看的是大纲级别：正文样式的段落直接设了 `w:outlineLvl`，在导航窗格
    和目录里就是标题，也可能挂在标题的编号列表上。不认的话整节并进上一节正文。

    只认段落上**直接**设的那个，不沿样式链往上找：目录标题之类的样式继承自
    Heading 1，会凭空多出一章。`outlineLvl=9` 是「正文文本」。大纲级别也会随复制
    粘贴被带到正文段落上，所以读起来是一句话的不算。
    """
    match = _HEADING.match(paragraph.style.name or "")
    if match:
        return int(match.group(1))
    ppr = paragraph._p.pPr
    outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    if outline is None:
        return None
    value = outline.get(qn("w:val"))
    if value is None or not value.isdigit() or int(value) > 8:
        return None
    if _reads_as_sentence(_paragraph_text(paragraph).strip()):
        return None
    return int(value) + 1


def _parse_date(raw: str) -> date | None:
    for date_format in _DATE_FORMATS:
        try:
            # 只取日期部分，时区无关
            return datetime.strptime(raw.strip(), date_format).date()  # noqa: DTZ007
        except ValueError:
            continue
    return None


def _table_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [
            "\n".join(_paragraph_text(p) for p in cell.paragraphs).strip()
            for cell in row.cells
        ]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _numbered_heading(text: str, dotted_are_headings: bool):
    """没有 Heading 样式时，从编号认标题。返回 (level, number, title) 或 None。"""
    found = parse_label(text)
    if found is None:
        return None
    if not dotted_are_headings and found.trailing_dot and len(found.parts) == 1:
        return None
    return len(found.parts), ".".join(str(part) for part in found.parts), found.title


# 标题总是短的。80 字符是形状判断：再长就是强调段落，不是标题。
_HEADING_MAX_CHARS = 80


# 缺了章号的编号：`.6.1 Title`。形状判断，换哪家机构都成立。
_MISSING_CHAPTER = re.compile(r"^\.\d+(?:\.\d+)*\s+\S")


def _reads_as_sentence(text: str) -> bool:
    """长、或以句读收尾——是一句话，不是标题。"""
    return len(text) > _HEADING_MAX_CHARS or text.rstrip().endswith((".", ";", ","))


def _body_point_size(paragraphs) -> float | None:
    """这份文档的正文字号——取出现次数最多的那个。"""
    sizes: dict[float, int] = {}
    for paragraph in paragraphs:
        runs = [run for run in paragraph.runs if run.text.strip()]
        points = [run.font.size.pt for run in runs if run.font.size]
        if points:
            sizes[max(points)] = sizes.get(max(points), 0) + 1
    return max(sizes, key=lambda size: sizes[size]) if sizes else None


def _looks_like_heading(paragraph, body_point: float | None) -> bool:
    """既无样式也无编号时的兜底：整段加粗、或字号比正文大，且短。

    判据一律**相对于这份文档**——写死 12pt / 14pt 就又变成一份文档的特征了。
    """
    text = _paragraph_text(paragraph).strip()
    if not text or len(text) > _HEADING_MAX_CHARS:
        return False
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    if all(run.bold for run in runs):
        return True
    points = [run.font.size.pt for run in runs if run.font.size]
    return bool(points) and body_point is not None and max(points) > body_point


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise ParseError(
                f"无法打开 docx 文件：{type(exc).__name__}",
                code="parse_error.docx_unreadable", params={"error": type(exc).__name__},
            ) from exc

        warnings: list[str] = []
        # Word 的标题编号不在文字里，是渲染时按 numbering.xml 画上去的。
        # 不还原的话 citation_label 只能退化成标题路径，而审计引用要的是 3.4.2。
        numbering = HeadingNumbering(document)
        # 样式在每个段落上优先；同一文档中的 Normal 编号标题仍可被识别。
        styled = any(_heading_level(paragraph) is not None for paragraph in document.paragraphs)
        survey = [
            found.parts
            for paragraph in document.paragraphs
            if (found := parse_label(_paragraph_text(paragraph).strip()))
        ]
        dotted_are_headings = acts_as_headings(survey)

        # 三级取舍：样式 > 编号 > 格式。样式最可靠，编号能给出层级，
        # 格式只说得出"这是个标题"，给不出层级也给不出编号——所以只在
        # 前两者都无所获时才用它。
        by_format = False
        body_point = None
        if not styled:
            numbered_candidates = sum(
                1
                for paragraph in document.paragraphs
                if _noise(_paragraph_text(paragraph).strip()) is None
                and _numbered_heading(_paragraph_text(paragraph).strip(), dotted_are_headings)
            )
            if numbered_candidates == 0:
                by_format = True
                body_point = _body_point_size(document.paragraphs)
        numbered = 0
        roots: list[ClauseNode] = []
        stack: list[ClauseNode] = []
        meta_fields: dict[str, str] = {}
        doc_type: str | None = None
        empty_headings = 0
        suspect_numbers = 0
        table_index = 0
        source_blocks: list[SourceBlock] = []

        def attach(node: ClauseNode, level: int) -> None:
            while stack and stack[-1].level >= level:
                stack.pop()
            if stack:
                stack[-1].children.append(node)
            else:
                roots.append(node)

        # **按文档顺序**遍历 body。`document.paragraphs` 不含表格，
        # `document.tables` 又丢掉位置信息——两个循环各走各的，表格在第几节里
        # 就永远看不出来了。实测一份技术标准因此有六个章节正文为空，
        # 它们的内容是文末孤零零的六张表。
        for body_index, element in enumerate(document.element.body.iterchildren(), start=1):
            if element.tag == qn("w:tbl"):
                table = Table(element, document)
                body = _table_text(table)
                if not body:
                    continue
                table_index += 1
                # 合并单元格在网格中重复出现：原文只存一次，同时保留所有覆盖位置。
                cells: dict[object, tuple[str, str, list[dict[str, int]]]] = {}
                for row_index, row in enumerate(table.rows, start=1):
                    for column, cell in enumerate(row.cells, start=1):
                        if cell._tc not in cells:
                            key = f"body:{body_index}:cell:{row_index}:{column}"
                            value = "\n".join(_paragraph_text(p) for p in cell.paragraphs)
                            cells[cell._tc] = (key, value, [])
                        cells[cell._tc][2].append({"row": row_index, "column": column})
                table_keys = []
                for key, value, positions in cells.values():
                    table_keys.append(key)
                    source_blocks.append(SourceBlock(
                        key=key, text=value, kind="table_cell",
                        locator={"body_index": body_index, "table": table_index,
                                 "positions": positions},
                    ))
                node = ClauseNode(
                    heading=f"Table {table_index}",
                    text=body,
                    level=(stack[-1].level + 1) if stack else 1,
                    kind="table",
                    source_keys=table_keys,
                )
                if stack:
                    stack[-1].children.append(node)
                else:
                    # 封面上的版本历史表、缩略语表不属于任何章节。
                    roots.append(node)
                continue
            if element.tag != qn("w:p"):
                continue

            paragraph = Paragraph(element, document)
            original_text = _paragraph_text(paragraph)
            text = original_text.strip()
            source_key = f"body:{body_index}:paragraph"
            if text:
                source_blocks.append(SourceBlock(
                    key=source_key, text=original_text, kind="paragraph",
                    locator={"body_index": body_index, "style": paragraph.style.name},
                ))
            if _mark_deleted(paragraph):
                # 不当标题、不推进编号；还剩的文字按正文留下。Word 里它会并到下一段
                # 开头，这里挂在当前章节末尾——位置差一段，但一个字不丢。
                if text:
                    if stack:
                        stack[-1].text = f"{stack[-1].text}\n{text}".strip()
                        stack[-1].source_keys.append(source_key)
                    else:
                        append_text(roots, text, source_key=source_key)
                continue
            level = _heading_level(paragraph)
            from_number = (
                None
                if level is not None or _noise(text) is not None
                else _numbered_heading(text, dotted_are_headings)
            )
            if styled and from_number is not None:
                label = parse_label(text)
                # 混合文档中的单层列表仍按正文保留，不能因 1/2/3 连续而升为章节。
                if label is not None and len(label.parts) == 1 and label.kind != "article":
                    has_children = any(parts[:-1] == label.parts for parts in survey)
                    if not has_children and not _looks_like_heading(paragraph, None):
                        from_number = None
                # 正文样式里以 "Section 5" 开头的完整句子是引用，不是标题。
                elif label is not None and label.kind == "article" and _reads_as_sentence(text):
                    from_number = None
                # 有样式的文档里，正文手打的多级编号只有**接得上**已打开的标题时
                # 才算下级标题：`6.1.2` 在 `6.1` 下面成立；第 6 章里一句
                # "2.1 of the Standard applies…" 是引用，当成标题会凭空补出第 2 章。
                # 「接得上」= 父级已打开，或与某个已打开的编号同父（`1.1` 之后的 `1.2`，
                # 标题本身不带编号时就是这样一路接下去的）。
                # 栈里还没有任何编号时（标题本身不带编号）无从判断，保持原样。
                elif label is not None and len(label.parts) > 1:
                    open_numbers = {node.number for node in stack if node.number}
                    reachable = open_numbers | {
                        number.rsplit(".", 1)[0] for number in open_numbers if "." in number
                    }
                    parent = from_number[1].rsplit(".", 1)[0]
                    if open_numbers and parent not in reachable:
                        from_number = None
            if styled and from_number is None and level is None and _MISSING_CHAPTER.match(text):
                # `.6.1` 这类缺了章号的编号只告警、不猜：猜错的条款号比没有更糟。
                suspect_numbers += 1
            if from_number is not None:
                level, number, title = from_number
                numbered += 1
                # 混合样式：未编号的真实父标题可接住其下的编号子标题。
                # 有编号的父级仍按编号血统处理，避免把 4.1 挂到 3 下。
                unnumbered_parent = next((
                    node for node in reversed(stack)
                    if node.level < level and node.number is None
                ), None)
                if unnumbered_parent is not None and all(
                    node.number is None or number.startswith(f"{node.number}.")
                    or node.level >= level for node in stack
                ):
                    node = ClauseNode(heading=title, text="", level=level, number=number,
                                      source_keys=[source_key])
                    attach(node, level)
                    stack.append(node)
                    continue
                # 按**编号血统**收栈，不是按层级。层级对得上、编号对不上时，
                # `4.1` 会挂到上一个一级节点 `3` 下面去——审计引用顺着它走
                # 会走到别的章节。实测 01/02/06/10 四份都栽在这里。
                #
                # 缺的祖先要补出来，理由同 numbering.py：不补的话 `4.1` 自己
                # 变成顶层条款，一份文档的"顶层"里混着二级三级编号。
                #
                # 先退到最近的真祖先，再只补它和本条之间缺的那几级。原先逐级比对时
                # 方向写反了：处理 `4.1.1` 的第一级祖先 `4` 时，栈顶的 `4.1` 不等于
                # `4`、`4` 也不以 `4.1.` 开头，于是真的 `4.1` 被弹掉，每个三级编号
                # 都各自挂在一个补出来的空 `4.1` 下面，父节点的正文进不了子条款的上下文。
                while stack and not (
                    stack[-1].number and number.startswith(f"{stack[-1].number}.")
                ):
                    stack.pop()
                parts = number.split(".")
                have = len(stack[-1].number.split(".")) if stack else 0
                for depth in range(have + 1, len(parts)):
                    ancestor = ".".join(parts[:depth])
                    placeholder = ClauseNode(
                        heading=ancestor, text="", level=depth, number=ancestor
                    )
                    (stack[-1].children if stack else roots).append(placeholder)
                    stack.append(placeholder)
                node = ClauseNode(heading=title, text="", level=level, number=number,
                                  source_keys=[source_key])
                (stack[-1].children if stack else roots).append(node)
                stack.append(node)
                continue

            if by_format and _looks_like_heading(paragraph, body_point):
                node = ClauseNode(heading=text, text="", level=1, kind="section",
                                  source_keys=[source_key])
                stack.clear()
                roots.append(node)
                stack.append(node)
                continue

            if level is not None:
                # 计数器必须按文档顺序推进，空标题也要走一遍，否则后面全错位。
                number = numbering.number_for(paragraph)
                # 显式写进标题的编号不在 numbering.xml 中。
                explicit = parse_label(text)
                if number is None and explicit is not None:
                    number = ".".join(str(part) for part in explicit.parts)
                    text = explicit.title
                if not text:
                    empty_headings += 1
                    continue
                if number:
                    numbered += 1
                node = ClauseNode(heading=text, text="", level=level, number=number,
                                  source_keys=[source_key])
                attach(node, level)
                stack.append(node)
                continue

            if not text:
                continue

            if match := _META.match(text):
                meta_fields.setdefault(match.group(1), match.group(2).strip())
            if doc_type is None and (match := _DOC_TYPE.search(text)):
                doc_type = match.group(1).lower()

            if stack:
                stack[-1].text = f"{stack[-1].text}\n{text}".strip()
                stack[-1].source_keys.append(source_key)
            else:
                append_text(roots, text, source_key=source_key)

        if empty_headings:
            warnings.append(
                ParseWarning(
                    f"已跳过 {empty_headings} 个空标题",
                    code="parse.empty_headings_skipped",
                    params={"count": empty_headings},
                )
            )
        if suspect_numbers:
            warnings.append(
                ParseWarning(
                    f"{suspect_numbers} 个疑似标题的编号缺了章号（如 .6.1），已按正文保留",
                    code="parse.suspect_numbering",
                    params={"count": suspect_numbers},
                )
            )
        if not any(node.kind == "section" for node in roots):
            warnings.append(
                ParseWarning(
                    "未找到任何标题样式，该文档可能不是标准 house style",
                    code="parse.no_heading_styles",
                )
            )
        if not numbered and any(node.kind == "section" for node in roots):
            # 说出来而不是静默退化：条款号会变成标题路径，审计引用会不好用。
            warnings.append(
                ParseWarning(
                    "未能还原任何标题编号，条款号将退化为标题路径",
                    code="parse.no_numbering_restored",
                )
            )

        meta = DocumentMeta(
            title=None,
            version=meta_fields.get("Version"),
            owner=meta_fields.get("Owner"),
            approver=meta_fields.get("Approver"),
            approved_date=_parse_date(meta_fields.get("Approved Date", "")),
            effective_date=_parse_date(meta_fields.get("Effective Date", "")),
            doc_type=doc_type,
        )
        # 同 PDF 路："标题"其实是段落前半句时，整段收回正文。
        demote_fragment_headings(roots)
        warnings.extend(structure_warnings(roots))
        return ParsedDocument(meta=meta, clauses=roots, warnings=warnings,
                              source_blocks=source_blocks, source_origin="docx_xml")
