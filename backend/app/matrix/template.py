"""有界的 XLSX 解析与纯校验；**从不求值表格公式**。"""

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

CANONICAL_FIELDS: tuple[str, ...] = (
    "code", "title", "statement", "category", "owner", "framework_refs", "note",
)
REQUIRED_FIELDS: tuple[str, ...] = ("title", "statement")
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 128
MAX_CELL_CHARS = 32_767
MAX_TEXT_CHARS = 8 * 1024 * 1024


class SheetRejected(ValueError):
    """带 code 的拒绝。路由层把 code 与 params 原样交给前端查词条，句子只是兜底。

    仍然是 ValueError：导入链路上层层 `except ValueError` 的地方不必改。
    """

    def __init__(self, message: str, *, code: str, params: dict[str, object] | None = None):
        super().__init__(message)
        self.code = code
        self.params = params or {}


@dataclass(frozen=True)
class Issue:
    """校验报告里的一条问题：前端按 code 组句，detail 是中文兜底。"""

    code: str
    detail: str
    params: dict[str, object] = field(default_factory=dict)

    def rejected(self) -> SheetRejected:
        return SheetRejected(self.detail, code=self.code, params=self.params)


@dataclass(frozen=True)
class Sheet:
    headers: list[str]
    rows: list[list[str]]


def _inspect_archive(content: bytes) -> None:
    if not content or len(content) > MAX_FILE_BYTES:
        raise SheetRejected("Excel 文件为空或超过 10 MiB 限制", code="matrix.file_empty_or_too_large", params={"limit_mib": MAX_FILE_BYTES // (1024 * 1024)})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 256 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
            raise SheetRejected("Excel 解压大小或文件数量超过限制", code="matrix.archive_too_large")
        if len({e.filename for e in entries}) != len(entries):
            raise SheetRejected("Excel 包含重复文件", code="matrix.archive_duplicate_entries")
        for entry in entries:
            if entry.flag_bits & 1 or entry.filename.lower().endswith(".bin"):
                raise SheetRejected("Excel 不支持加密文件或宏", code="matrix.encrypted_or_macro")
            if not entry.filename.endswith((".xml", ".rels")):
                continue
            raw = archive.read(entry)
            # XLSX 里的 XML 是 UTF-8。在任何一个 XML 解析器看到不可信声明之前，
            # 先挡掉 DTD／实体声明——连 UTF-16 的写法一起挡。
            safe_scan = raw.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in safe_scan or b"<!ENTITY" in safe_scan:
                raise SheetRejected("Excel XML 不允许实体或 DTD", code="matrix.xml_dtd")
            if not entry.filename.startswith("xl/worksheets/"):
                continue
            for _, element in ElementTree.iterparse(io.BytesIO(raw), events=("start",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "row" and int(element.get("r", "1")) > MAX_ROWS + 1:
                    raise SheetRejected("Excel 行数超过限制", code="matrix.sheet_too_large", params={"rows": MAX_ROWS, "columns": MAX_COLUMNS})
                if tag == "c":
                    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", element.get("r", ""))
                    if not match:
                        raise SheetRejected("Excel 单元格坐标无效", code="matrix.cell_ref_invalid")
                    if (column_index_from_string(match[1]) > MAX_COLUMNS
                            or int(match[2]) > MAX_ROWS + 1):
                        raise SheetRejected("Excel 行列数超过限制", code="matrix.sheet_too_large", params={"rows": MAX_ROWS, "columns": MAX_COLUMNS})
                element.clear()


def read_sheet(content: bytes, *, max_rows: int | None = None) -> Sheet:
    if max_rows is not None and (type(max_rows) is not int or max_rows < 0):
        raise ValueError("max_rows 必须为非负整数")
    book = None
    try:
        _inspect_archive(content)
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        worksheet = book.active
        if worksheet is None or not hasattr(worksheet, "iter_rows"):
            raise SheetRejected("Excel 缺少数据工作表", code="matrix.no_worksheet")
        # Do not trust attacker-controlled worksheet dimensions (e.g. XFD1048576).
        worksheet.reset_dimensions()
        headers: list[str] = []
        rows: list[list[str]] = []
        text_chars = 0
        for number, cells in enumerate(worksheet.iter_rows(), start=1):
            if number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                raise SheetRejected("Excel 行列数超过限制", code="matrix.sheet_too_large", params={"rows": MAX_ROWS, "columns": MAX_COLUMNS})
            row: list[str] = []
            for cell in cells:
                if cell.data_type in {"f", "e"}:
                    raise SheetRejected("Excel 包含公式或错误单元格，请先粘贴为值", code="matrix.formula_cells")
                value = "" if cell.value is None else str(cell.value).strip()
                if len(value) > MAX_CELL_CHARS:
                    raise SheetRejected("Excel 单元格文本超过限制", code="matrix.cell_too_long", params={"limit": MAX_CELL_CHARS})
                text_chars += len(value)
                if text_chars > MAX_TEXT_CHARS:
                    raise SheetRejected("Excel 文本总量超过限制", code="matrix.text_too_large")
                row.append(value)
            if number == 1:
                headers = row
                while headers and not headers[-1]:
                    headers.pop()
                if max_rows == 0:
                    break
            elif any(row):
                rows.append(row)
                if max_rows is not None and len(rows) >= max_rows:
                    break
        return Sheet(headers, rows)
    except ValueError:
        raise
    except Exception as exc:
        raise SheetRejected(
            f"无法读取 Excel 文件：{type(exc).__name__}",
            code="matrix.unreadable", params={"error": type(exc).__name__},
        ) from exc
    finally:
        if book is not None:
            book.close()


def validate(sheet: Sheet, mapping: dict[str, str]) -> list[Issue]:
    report: list[Issue] = []
    if not isinstance(mapping, dict) or not all(
        isinstance(k, str) and isinstance(v, str) for k, v in mapping.items()
    ):
        return [Issue(code="matrix.mapping_not_object", detail="mapping 必须为目标字段到源列名的字符串对象")]
    if not sheet.headers or any(not h for h in sheet.headers):
        report.append(Issue(code="matrix.header_blank", detail="Excel 表头为空或包含空列名"))
    if len(set(sheet.headers)) != len(sheet.headers):
        report.append(Issue(code="matrix.header_duplicate", detail="Excel 表头重复，无法确定源列"))
    for name in REQUIRED_FIELDS:
        if name not in mapping:
            report.append(Issue(
                code="matrix.required_field_unmapped",
                detail=f"必填字段 {name} 没有映射到任何一列", params={"field": name},
            ))
    index: dict[str, int] = {}
    for name, column in mapping.items():
        if name not in CANONICAL_FIELDS:
            report.append(Issue(
                code="matrix.unknown_field", detail=f"未知的目标字段 {name}", params={"field": name},
            ))
        elif column not in sheet.headers:
            report.append(Issue(
                code="matrix.column_missing",
                detail=f"映射到了不存在的列「{column}」", params={"column": column},
            ))
        else:
            index[name] = sheet.headers.index(column)
    if report:
        return report
    if not sheet.rows:
        report.append(Issue(code="matrix.no_rows", detail="Excel 没有数据行"))
    for name in REQUIRED_FIELDS:
        blanks = sum(not _cell(row, index[name]) for row in sheet.rows)
        if blanks:
            report.append(Issue(
                code="matrix.required_field_blank",
                detail=f"有 {blanks} 行的 {name} 为空", params={"count": blanks, "field": name},
            ))
    for name, limit in (("code", 64), ("title", 500), ("category", 100)):
        if name in index and any(len(_cell(r, index[name])) > limit for r in sheet.rows):
            report.append(Issue(
                code="matrix.field_too_long",
                detail=f"{name} 长度超过 {limit}", params={"field": name, "limit": limit},
            ))
    if any(any(row[len(sheet.headers):]) for row in sheet.rows):
        report.append(Issue(code="matrix.data_without_header", detail="存在没有表头的数据列"))
    if "code" in index:
        counts = Counter(_cell(row, index["code"]) for row in sheet.rows)
        duplicates = sorted(code for code, count in counts.items() if code and count > 1)
        if duplicates:
            report.append(Issue(
                code="matrix.duplicate_codes",
                detail=f"编号重复：{'、'.join(duplicates[:20])}", params={"codes": duplicates[:20]},
            ))
    return report


def _cell(row: list[str], position: int) -> str:
    return row[position].strip() if position < len(row) else ""
