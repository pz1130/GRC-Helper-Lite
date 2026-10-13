"""评估语料：全部来自仓库自带的公开数据，可复现。

· 映射：NIST CSF 2.0（seeds/nist-csf-2.0.csv）为框架；OLIR 夹具里出现的 800-53 Rev.5
  基础控制项（seeds/nist-800-53-r5.csv）充当「本机构控制点」，陈述截到 250 字符——
  与 KNOWN_ISSUES K16 的实测口径一致。正例 = OLIR 的 (CSF 子类, 800-53 基础控制) 对。
· 抽取：用 800-53 基础控制项的原文拼成一份 Word「制度」（编号章节 + 条款正文）。
  NIST 出版物是美国政府作品，公有领域。它不是真实机构制度，只用于回归对比。
"""

import csv
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SEEDS = ROOT / "seeds"
OLIR = ROOT / "tests" / "fixtures" / "olir_csf_800-53.json"
CONTROL_STATEMENT_CHARS = 250
EXTRACTION_FAMILIES = ("AC", "IA", "IR", "CP")
EXTRACTION_PER_FAMILY = 4


def _rows(name: str) -> list[dict[str, str]]:
    return list(csv.DictReader((SEEDS / name).read_text(encoding="utf-8-sig").splitlines()))


def olir_pairs() -> dict[str, set[str]]:
    data = json.loads(OLIR.read_text(encoding="utf-8"))
    return {pair["csf"]: set(pair["controls"]) for pair in data["pairs"]}


def mapping_controls() -> list[dict[str, str]]:
    wanted = set().union(*olir_pairs().values())
    rows = {row["code"]: row for row in _rows("nist-800-53-r5.csv")}
    return [
        {"code": code, "title": rows[code]["title"],
         "statement": (rows[code]["description"] or "")[:CONTROL_STATEMENT_CHARS]}
        for code in sorted(wanted) if code in rows
    ]


def extraction_sections() -> list[tuple[str, str, list[tuple[str, str]]]]:
    rows = _rows("nist-800-53-r5.csv")
    families = {row["code"]: row["title"] for row in rows if not row["parent_code"]}
    out = []
    for family in EXTRACTION_FAMILIES:
        controls = [
            (row["title"], row["description"]) for row in rows
            if row["parent_code"] == family and row["description"]
            and "withdrawn" not in row["description"].casefold()
        ][:EXTRACTION_PER_FAMILY]
        out.append((family, families[family], controls))
    return out


def build_extraction_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Information Security Policy (evaluation corpus from NIST SP 800-53 Rev.5)", 0)
    for n, (_, title, controls) in enumerate(extraction_sections(), start=1):
        doc.add_heading(f"{n} {title}", 1)
        for m, (ctitle, text) in enumerate(controls, start=1):
            doc.add_heading(f"{n}.{m} {ctitle}", 2)
            doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
