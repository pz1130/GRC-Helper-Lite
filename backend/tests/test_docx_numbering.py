"""Word 自动编号的还原。

Word 的标题编号不在文字里：`w:t` 只有 "Introduction"，前面那个 "1" 是渲染时按
numbering.xml 画上去的。不还原就只能让 citation_label 退化成标题路径，而审计
引用要的是 3.4.2。
"""

import re
from pathlib import Path

import pytest
from docx import Document

from app.parsing.docx_numbering import HeadingNumbering, format_counter


def test_decimal_letters_and_roman():
    assert format_counter(3, "decimal") == "3"
    assert format_counter(1, "lowerLetter") == "a"
    assert format_counter(27, "lowerLetter") == "aa"
    assert format_counter(4, "upperLetter") == "D"
    assert format_counter(9, "lowerRoman") == "ix"
    assert format_counter(2024, "upperRoman") == "MMXXIV"


def test_an_unknown_format_returns_none_instead_of_guessing():
    """错的条款号比没有更糟——审计引用会指向不存在的位置。"""
    assert format_counter(1, "bullet") is None
    assert format_counter(1, "chineseCounting") is None


# ---- 状态机 ----


class _Stub:
    """最小的段落替身：只需要 _p（找 pPr）与 style.style_id。"""

    def __init__(self, root, style_id: str) -> None:
        self._p = root
        self.style = type("S", (), {"style_id": style_id})()


def _docx_with_numbering(tmp_path, levels: str, styles: str):
    """造一个带指定 numbering/styles 的 docx，用 python-docx 打开。"""
    import shutil
    import zipfile

    source = tmp_path / "base.docx"
    Document().save(source)
    target = tmp_path / "built.docx"
    shutil.copy(source, target)

    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    numbering = f'<w:numbering {W}>{levels}</w:numbering>'
    style_xml = f'<w:styles {W}>{styles}</w:styles>'
    with zipfile.ZipFile(source) as src:
        names = [n for n in src.namelist() if n not in
                 ("word/numbering.xml", "word/styles.xml")]
        parts = {n: src.read(n) for n in names}
        rels = src.read("word/_rels/document.xml.rels").decode()
    if "numbering.xml" not in rels:
        rels = rels.replace("</Relationships>",
            '<Relationship Id="rIdNum" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/numbering" Target="numbering.xml"/>'
            "</Relationships>")
    parts["word/_rels/document.xml.rels"] = rels.encode()
    with zipfile.ZipFile(target, "w") as out:
        for name, data in parts.items():
            out.writestr(name, data)
        out.writestr("word/numbering.xml", numbering)
        out.writestr("word/styles.xml", style_xml)
    return Document(str(target))


LEVELS = "".join(
    f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    f'<w:lvlText w:val="{".".join("%" + str(k + 1) for k in range(i + 1))}"/></w:lvl>'
    for i in range(4)
)
NUMBERING = (
    f'<w:abstractNum w:abstractNumId="15">{LEVELS}</w:abstractNum>'
    '<w:num w:numId="3"><w:abstractNumId w:val="15"/></w:num>'
)
# Heading1 只声明 ilvl，numId 要沿 basedOn 链从 Heading2 继承——真实文档就是这样。
STYLES = (
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:basedOn w:val="Heading2"/>'
    '<w:pPr><w:numPr><w:ilvl w:val="0"/></w:numPr></w:pPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading2">'
    '<w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="3"/></w:numPr></w:pPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading3"><w:basedOn w:val="Heading2"/>'
    '<w:pPr><w:numPr><w:ilvl w:val="2"/></w:numPr></w:pPr></w:style>'
)


def _sequence(document, plan: list[str]) -> list[str | None]:
    from docx.oxml.ns import qn

    numbering = HeadingNumbering(document)
    out = []
    for style_id in plan:
        paragraph = document.add_paragraph()
        paragraph._p.get_or_add_pPr()
        out.append(numbering.number_for(_Stub(paragraph._p, style_id)))
        assert paragraph._p.find(qn("w:pPr")) is not None
    return out


def test_numbers_nest_and_reset(tmp_path):
    document = _docx_with_numbering(tmp_path, NUMBERING, STYLES)
    got = _sequence(document, [
        "Heading1", "Heading2", "Heading2", "Heading1", "Heading2", "Heading3", "Heading3",
        "Heading2", "Heading1",
    ])
    assert got == ["1", "1.1", "1.2", "2", "2.1", "2.1.1", "2.1.2", "2.2", "3"]


def test_num_id_is_inherited_through_the_based_on_chain(tmp_path):
    """Heading1 自己没有 numId，只有 ilvl——真实文档就是这样写的。"""
    document = _docx_with_numbering(tmp_path, NUMBERING, STYLES)
    assert _sequence(document, ["Heading1"]) == ["1"]


def test_a_paragraph_with_no_numbering_gets_nothing(tmp_path):
    document = _docx_with_numbering(tmp_path, NUMBERING, STYLES)
    assert _sequence(document, ["BodyText"]) == [None]


def test_a_document_without_numbering_part_does_not_explode(tmp_path):
    """没有 numbering.xml 的 docx 要安静地返回 None，而不是抛异常。"""
    document = Document()
    numbering = HeadingNumbering(document)
    paragraph = document.add_paragraph()
    paragraph._p.get_or_add_pPr()
    assert numbering.number_for(_Stub(paragraph._p, "Heading1")) is None


# ---- 真实文档 ----

# 按目录里现有的 .docx 取，而不是写死文件名：语料换一批这条仍然跑得起来，
# 也不必把哪家机构的文件名留在版本库里。
_DOCX = sorted(Path("/samples").glob("*.docx")) if Path("/samples").is_dir() else []
REAL = str(_DOCX[0]) if _DOCX else ""
EXPECTED_HEAD = ["1", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "2", "2.1", "2.2", "2.3"]


@pytest.mark.skipif(not _DOCX, reason="需挂载 sample docs，见 make corpus")
def test_the_real_document_numbers_match_its_own_table_of_contents():
    """未改动的章节与文档自带目录逐条一致。

    目录本身是**过期的**：文档后来插入了一个一级标题并把原来的降了一级，目录没
    重新生成。所以只比对目录仍然覆盖的部分——这也正是当初放弃「直接读目录」
    方案的理由，那样会在这份文档上算出错的编号。
    """
    document = Document(REAL)
    numbering = HeadingNumbering(document)
    heading = re.compile(r"^Heading (\d+)$")
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    got = []
    for paragraph in document.paragraphs:
        if not heading.match(paragraph.style.name or ""):
            continue
        number = numbering.number_for(paragraph)
        text = "".join(t.text or "" for t in paragraph._p.iter(f"{W}t")).strip()
        if text:
            got.append(number)

    assert got[:len(EXPECTED_HEAD)] == EXPECTED_HEAD
    assert got[-2:] == ["4", "5"], "尾部两节未改动，编号应当仍是 4 与 5"
    assert all(n is not None for n in got), "32 条标题此前一个编号都没有"


# ---- 修订痕迹 ----
#
# 定稿前的制度常带着成百处修订。一个「段落标记已删除」的空标题段落，在 Word 的
# 最终显示里并入下一段、根本不存在，也不占编号；解析器却照样推进了计数器，
# 后面整节编号错位（实测一份真实文档的 2.1 / 2.2 / 2.3 被解析成 2.3 / 2.4 / 2.5）。
# 条款号错了，审计引用就指到别处。

NAMED_STYLES = STYLES.replace(
    'w:styleId="Heading1">', 'w:styleId="Heading1"><w:name w:val="heading 1"/>'
).replace(
    'w:styleId="Heading2">', 'w:styleId="Heading2"><w:name w:val="heading 2"/>'
).replace(
    'w:styleId="Heading3">', 'w:styleId="Heading3"><w:name w:val="heading 3"/>'
)
_DELETED_MARK = ('<w:rPr><w:del w:id="{id}" w:author="Reviewer" '
                 'w:date="2025-06-01T00:00:00Z"/></w:rPr>')


def _parse_headings(tmp_path, plan: list[tuple[str, str, bool]]):
    """plan: (styleId, 文字, 段落标记是否已被修订删除)。返回 (编号, 标题) 列表与正文。"""
    from lxml import etree

    from app.parsing.docx_parser import DocxParser

    document = _docx_with_numbering(tmp_path, NUMBERING, NAMED_STYLES)
    body = document.element.body
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for n, (style, text, deleted) in enumerate(plan):
        mark = _DELETED_MARK.format(id=n + 1) if deleted else ""
        run = f"<w:r><w:t>{text}</w:t></w:r>" if text else ""
        paragraph = etree.fromstring(
            f'<w:p xmlns:w="{W}"><w:pPr><w:pStyle w:val="{style}"/>{mark}</w:pPr>{run}</w:p>'
        )
        body.insert(len(body) - 1, paragraph)  # 保持 sectPr 在最后
    path = tmp_path / "tracked.docx"
    document.save(str(path))
    parsed = DocxParser().parse(path)

    def walk(nodes):
        for node in nodes:
            yield node
            yield from walk(node.children)

    nodes = [n for n in walk(parsed.clauses) if n.kind == "section"]
    return [(n.number, n.heading) for n in nodes], nodes, parsed


def test_an_empty_heading_deleted_in_track_changes_does_not_take_a_number(tmp_path):
    headings, _, _ = _parse_headings(tmp_path, [
        ("Heading1", "Introduction", False),
        ("Heading2", "Purpose", False),
        ("Heading1", "Roles and Responsibilities", False),
        ("Heading2", "", True),
        ("Heading2", "", True),
        ("Heading2", "Approver", False),
        ("Heading2", "Owner", False),
        ("Heading1", "Process", False),
    ])
    assert headings == [
        ("1", "Introduction"), ("1.1", "Purpose"),
        ("2", "Roles and Responsibilities"), ("2.1", "Approver"), ("2.2", "Owner"),
        ("3", "Process"),
    ]


def test_an_ordinary_empty_heading_still_takes_its_number(tmp_path):
    """没被删除的空标题在 Word 里照样显示编号——这一条不能跟着改。"""
    headings, _, _ = _parse_headings(tmp_path, [
        ("Heading1", "Roles and Responsibilities", False),
        ("Heading2", "", False),
        ("Heading2", "Approver", False),
    ])
    assert headings == [("1", "Roles and Responsibilities"), ("1.2", "Approver")]


def test_a_heading_whose_mark_was_deleted_keeps_its_words_as_body_text(tmp_path):
    """标记删了、文字还在：Word 里这段并进下一段。不当标题、不占编号，但字不能丢。"""
    headings, nodes, parsed = _parse_headings(tmp_path, [
        ("Heading1", "Roles", False),
        ("Heading2", "Leftover words", True),
        ("Heading2", "Approver", False),
    ])
    assert headings == [("1", "Roles"), ("1.1", "Approver")]
    assert any("Leftover words" in (n.text or "") for n in nodes)
    assert not any("跳过" in w and "空标题" in w for w in parsed.warnings)


def test_a_trailing_dot_in_the_level_text_is_not_part_of_the_clause_number(tmp_path):
    """`lvlText="%1."` 在 Word 里画成「1.」，条款号却是 1——带着点的话 `1.1` 接不上 `1.`。"""
    dotted = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        f'<w:lvlText w:val="{".".join("%" + str(k + 1) for k in range(i + 1))}."/></w:lvl>'
        for i in range(3)
    )
    numbering = (
        f'<w:abstractNum w:abstractNumId="15">{dotted}</w:abstractNum>'
        '<w:num w:numId="3"><w:abstractNumId w:val="15"/></w:num>'
    )
    document = _docx_with_numbering(tmp_path, numbering, STYLES)
    assert _sequence(document, ["Heading1", "Heading2", "Heading1"]) == ["1", "1.1", "2"]


# ---- 大纲级别 ----
#
# Word 认标题看的是「大纲级别」，不只是 Heading 样式：正文样式的段落直接设了
# `w:outlineLvl`，在导航窗格和目录里就是一个标题，常常还挂在标题的编号列表上。
# 若忽略大纲级别，章节会并进上一节正文，附录的子节也会挂错层级。

_STYLES_WITH_NORMAL = NAMED_STYLES + (
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
    '<w:name w:val="Normal"/></w:style>'
)


def _parse_paragraphs(tmp_path, plan: list[tuple[str, str, str]]):
    """plan: (styleId, 文字, 额外的 pPr 内容)。返回 (编号, 标题) 列表与解析结果。"""
    from lxml import etree

    from app.parsing.docx_parser import DocxParser

    document = _docx_with_numbering(tmp_path, NUMBERING, _STYLES_WITH_NORMAL)
    body = document.element.body
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for style, text, extra in plan:
        paragraph = etree.fromstring(
            f'<w:p xmlns:w="{W}"><w:pPr><w:pStyle w:val="{style}"/>{extra}</w:pPr>'
            f"<w:r><w:t>{text}</w:t></w:r></w:p>"
        )
        body.insert(len(body) - 1, paragraph)
    path = tmp_path / "outline.docx"
    document.save(str(path))
    parsed = DocxParser().parse(path)

    def walk(nodes):
        for node in nodes:
            yield node
            yield from walk(node.children)

    return [(n.number, n.heading) for n in walk(parsed.clauses) if n.kind == "section"], parsed


_TOP_LEVEL = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr><w:outlineLvl w:val="0"/>'


def test_a_body_style_paragraph_with_an_outline_level_is_a_heading(tmp_path):
    headings, parsed = _parse_paragraphs(tmp_path, [
        ("Heading1", "Introduction", ""),
        ("Normal", "This procedure applies to all teams.", ""),
        ("Heading1", "Document Archive", ""),
        ("Normal", "Records are kept in the archive.", ""),
        ("Normal", "Exceptions", _TOP_LEVEL),
        ("Normal", "Exceptions must be approved by the owner.", ""),
    ])
    assert headings == [("1", "Introduction"), ("2", "Document Archive"), ("3", "Exceptions")]
    archive, exceptions = parsed.clauses[1], parsed.clauses[2]
    assert archive.text == "Records are kept in the archive."
    assert exceptions.text == "Exceptions must be approved by the owner."


def test_an_unnumbered_outline_heading_takes_its_subsections_with_it(tmp_path):
    """附录标题不在编号列表上：它自己不占号，但它下面的小节不能挂到前一章去。"""
    _, parsed = _parse_paragraphs(tmp_path, [
        ("Heading1", "Exceptions", ""),
        ("Normal", "Exceptions must be approved by the owner.", ""),
        ("Normal", "Appendix A – Tool Steps", '<w:outlineLvl w:val="0"/>'),
        ("Heading2", "Add accounts", ""),
        ("Normal", "Owners add accounts in the vault.", ""),
    ])
    exceptions, appendix = parsed.clauses
    assert exceptions.children == []
    assert (appendix.number, appendix.heading) == (None, "Appendix A – Tool Steps")
    assert [c.heading for c in appendix.children] == ["Add accounts"]


def test_a_sentence_carrying_an_outline_level_stays_in_the_body(tmp_path):
    """大纲级别也会被误带到正文段落上（复制粘贴来的格式）。一句话不是标题。"""
    headings, parsed = _parse_paragraphs(tmp_path, [
        ("Heading1", "Access", ""),
        ("Normal", "Teams must review access every quarter and keep the evidence.",
         '<w:outlineLvl w:val="0"/>'),
    ])
    assert headings == [("1", "Access")]
    assert "every quarter" in parsed.clauses[0].text


def test_outline_level_nine_is_body_text(tmp_path):
    """`outlineLvl=9` 是 Word 的「正文文本」，不是第十级标题。"""
    headings, _ = _parse_paragraphs(tmp_path, [
        ("Heading1", "Access", ""),
        ("Normal", "Quarterly review", '<w:outlineLvl w:val="9"/>'),
    ])
    assert headings == [("1", "Access")]
