from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from sqlalchemy import select

from app.audit.citations import AnswerCitationValidator
from app.audit.history import similar_history, similarity
from app.audit.models import (
    AnswerDraft,
    AuditEngagement,
    AuditType,
    EngagementStatus,
    QuestionStatus,
)
from app.audit.service import finalize_answer, generate_answer, import_questions
from app.clauses.models import Clause
from app.controls.models import Control
from app.environment.models import (
    HowEnforced,
    Implementation,
    ImplementationStatus,
    TechAsset,
    TechAssetCategory,
    TechAssetEnvironment,
    TechAssetStatus,
)
from app.errors import Conflict
from app.evidence.models import EvidenceCadence, EvidenceItem, EvidenceStatus, EvidenceType
from app.frameworks.models import Framework, FrameworkItem, Mapping, MappingStrength
from app.iam.models import AuditLog, User
from app.iam.permissions import Role
from app.iam.security import create_access_token, hash_password
from app.ingest.models import DocStatus, DocType, Document
from app.llm.runner import ValidatedResult
from app.review.models import Proposal, ProposalKind, ProposalStatus
from app.review.service import Decision, decide
from app.search.schemas import SearchHit, SearchResponse


async def _user(db_session, role=Role.GRC_LEAD):
    user = User(
        email=f"{role.value}@example.com", name="User", role=role, password_hash=hash_password("pw")
    )
    db_session.add(user)
    await db_session.flush()
    return user


def _headers(user):
    return {"Authorization": f"Bearer {create_access_token(user.id, user.role)}"}


async def _engagement(db_session, user):
    row = AuditEngagement(
        name="ISO audit",
        audit_type=AuditType.EXTERNAL,
        status=EngagementStatus.PREPARING,
        created_by=user.id,
    )
    db_session.add(row)
    await db_session.flush()
    return row


@pytest.mark.asyncio
async def test_import_questions_preserves_order_and_audits(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)

    rows = await import_questions(
        db_session, engagement.id, "First question?\n- Second question?", "en", actor=actor
    )

    assert [(row.seq, row.question_text) for row in rows] == [
        (1, "First question?"),
        (2, "Second question?"),
    ]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_questions.import")
    )
    assert audit.after["count"] == 2


@pytest.mark.asyncio
async def test_accepting_answer_proposal_materializes_a_draft(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (
        await import_questions(
            db_session, engagement.id, "How is access reviewed?", "en", actor=actor
        )
    )[0]
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="a" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    control = Control(
        code="C-AUDIT", title="Access review", statement="Access is reviewed quarterly."
    )
    db_session.add_all([clause, control])
    await db_session.flush()
    payload = {
        "question_id": question.id,
        "body": "Access is reviewed quarterly.",
        "language": "en",
        "citations": [{"clause_id": clause.id, "quote": "Access is reviewed quarterly."}],
        "cited_control_ids": [control.id],
        "suggested_evidence_ids": [],
        "gap_notes": "No evidence is registered.",
        "confidence": 0.9,
    }
    proposal = Proposal(
        kind=ProposalKind.ANSWER,
        payload=payload,
        citations=payload["citations"],
        confidence=0.9,
        status=ProposalStatus.PENDING,
    )
    db_session.add(proposal)
    await db_session.flush()

    await decide(db_session, proposal.id, actor=actor, decision=Decision.ACCEPT)

    draft = await db_session.scalar(select(AnswerDraft))
    assert draft.body == payload["body"]
    assert draft.cited_clause_ids == [clause.id]
    assert question.status is QuestionStatus.DRAFTED


@pytest.mark.asyncio
async def test_finalize_moves_answer_into_history(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (await import_questions(db_session, engagement.id, "Question?", "en", actor=actor))[
        0
    ]
    answer = AnswerDraft(
        question_id=question.id,
        body="Final answer",
        language="en",
        cited_clause_ids=[],
        cited_control_ids=[],
        suggested_evidence_ids=[],
        gap_notes="",
        reviewed_by=actor.id,
    )
    db_session.add(answer)
    await db_session.flush()

    await finalize_answer(db_session, answer, actor=actor, seen_body=answer.body)
    response = await client.get("/api/audit/history", headers=_headers(actor))

    assert response.status_code == 200
    assert response.json()[0]["question_text"] == "Question?"
    assert response.json()[0]["final_body"] == "Final answer"
    assert question.status is QuestionStatus.FINALIZED


@pytest.mark.asyncio
async def test_viewer_cannot_create_engagement(client, db_session):
    viewer = await _user(db_session, Role.VIEWER)
    response = await client.post(
        "/api/audit/engagements",
        headers=_headers(viewer),
        json={"name": "Audit", "audit_type": "internal", "status": "preparing"},
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_create_engagement_with_audit_period(client, db_session):
    """带审计期间建项目。日期原样进审计日志的 JSONB 会序列化失败，整个请求 500。"""
    lead = await _user(db_session)
    response = await client.post(
        "/api/audit/engagements",
        headers=_headers(lead),
        json={
            "name": "Access audit",
            "audit_type": "internal",
            "period_start": "2026-01-01",
            "period_end": "2026-06-30",
        },
    )
    assert response.status_code == 201
    assert response.json()["period_start"] == "2026-01-01"

    entry = (
        await db_session.execute(
            select(AuditLog).where(AuditLog.action == "audit_engagement.create")
        )
    ).scalar_one()
    assert entry.after["period_start"] == "2026-01-01"
    assert entry.after["period_end"] == "2026-06-30"


@pytest.mark.asyncio
async def test_generation_creates_answer_proposal_not_formal_draft(db_session, monkeypatch):
    from app.audit import service

    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (
        await import_questions(
            db_session, engagement.id, "How is access reviewed?", "en", actor=actor
        )
    )[0]
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="b" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    db_session.add(clause)
    await db_session.flush()

    async def fake_search(*args, **kwargs):
        return SearchResponse(
            query=question.question_text,
            expanded_terms=[],
            vector_used=False,
            hits=[
                SearchHit(
                    chunk_id=1,
                    clause_id=clause.id,
                    document_id=document.id,
                    document_title=document.title,
                    citation_label="1",
                    heading_path="Review",
                    text=clause.text,
                    score=1,
                    rank_fulltext=1,
                    rank_vector=None,
                )
            ],
        )

    captured = {}

    async def fake_run(*args, **kwargs):
        captured.update(kwargs)
        return ValidatedResult(
            payload={
                "body": "Access is reviewed quarterly.",
                "language": "en",
                "citations": [{"clause_id": clause.id, "quote": "Access is reviewed quarterly."}],
                "cited_control_ids": [],
                "suggested_evidence_ids": [],
                "gap_notes": "",
                "confidence": 0.9,
            },
            confidence=0.9,
            llm_call_id=None,
        )

    monkeypatch.setattr(service, "search", fake_search)
    monkeypatch.setattr(service, "run", fake_run)
    proposal = await service.generate_answer(db_session, question.id, "en", actor=actor)

    assert proposal.kind is ProposalKind.ANSWER
    assert proposal.status is ProposalStatus.PENDING
    assert await db_session.scalar(select(AnswerDraft)) is None
    # 生成时的校验要拦下正文里的库内 clause_id，并告诉模型该写哪个条款号。
    leaked = await captured["citation_validator"].check(
        {
            "body": f"Access is reviewed quarterly (Clause {clause.id}).",
            "gap_notes": "",
            "citations": [{"clause_id": clause.id}],
            "cited_control_ids": [],
            "suggested_evidence_ids": [],
        }
    )
    assert leaked is not None and '"1"' in leaked


@pytest.mark.asyncio
async def test_answer_validator_rejects_control_outside_retrieval(db_session):
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="c" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    db_session.add(clause)
    await db_session.flush()
    validator = AnswerCitationValidator(
        db_session,
        allowed_clause_ids={clause.id},
        allowed_control_ids=set(),
        allowed_evidence_ids=set(),
    )

    reason = await validator.check(
        {
            "citations": [{"clause_id": clause.id, "quote": clause.text}],
            "cited_control_ids": [999],
            "suggested_evidence_ids": [],
        }
    )

    assert reason == "control 999 is not among the retrieved results"


async def _labelled_clause(db_session, label="4.2"):
    document = Document(
        title="Access Policy",
        doc_type=DocType.POLICY,
        file_hash="d" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number=label,
        heading="Multi-factor Authentication",
        heading_path="Multi-factor Authentication",
        citation_label=label,
        text="Multi-factor authentication shall be required for remote access.",
        order_index=0,
        level=2,
    )
    db_session.add(clause)
    await db_session.flush()
    return clause


def _answer_payload(clause_id, body, gap_notes=""):
    return {
        "body": body,
        "gap_notes": gap_notes,
        "citations": [{"clause_id": clause_id}],
        "cited_control_ids": [],
        "suggested_evidence_ids": [],
    }


@pytest.mark.parametrize("field", ["body", "gap_notes"])
@pytest.mark.asyncio
async def test_generated_answer_must_not_expose_internal_clause_ids(db_session, field):
    """提示词里给模型的是库内 clause_id；它写进正文，读答复的人看到「Clause 63」对不上任何条款号。"""
    clause = await _labelled_clause(db_session)
    validator = AnswerCitationValidator(
        db_session,
        allowed_clause_ids={clause.id},
        allowed_control_ids=set(),
        allowed_evidence_ids=set(),
        require_quotes=False,
        clause_labels={clause.id: "4.2"},
    )
    text = f"The requirement is stated (Clause {clause.id}, label 4.2)."
    payload = (
        _answer_payload(clause.id, "MFA is required.", text)
        if field == "gap_notes"
        else _answer_payload(clause.id, text)
    )

    reason = await validator.check(payload)

    assert reason == (
        f"the answer refers to internal clause ID {clause.id}; refer to the clause by its "
        'document title and clause number (for example "4.2") instead'
    )


@pytest.mark.asyncio
async def test_generated_answer_may_name_clause_numbers(db_session):
    clause = await _labelled_clause(db_session)
    validator = AnswerCitationValidator(
        db_session,
        allowed_clause_ids={clause.id},
        allowed_control_ids=set(),
        allowed_evidence_ids=set(),
        require_quotes=False,
        clause_labels={clause.id: "4.2"},
    )

    reason = await validator.check(
        _answer_payload(clause.id, "Access Policy clause 4.2 requires MFA for remote access.")
    )

    assert reason is None


@pytest.mark.asyncio
async def test_accepting_an_older_answer_does_not_apply_the_internal_id_check(db_session):
    """接受已有提案时不带 clause_labels：旧答复里的「Clause 63」由审核人改，不能让它无法接受。"""
    clause = await _labelled_clause(db_session)
    validator = AnswerCitationValidator(
        db_session, allowed_clause_ids={clause.id}, require_quotes=False
    )

    reason = await validator.check(_answer_payload(clause.id, f"See Clause {clause.id}."))

    assert reason is None


@pytest.mark.asyncio
async def test_generation_rejects_a_second_pending_answer(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (await import_questions(db_session, engagement.id, "Question?", "en", actor=actor))[
        0
    ]
    db_session.add(
        Proposal(
            kind=ProposalKind.ANSWER,
            payload={"question_id": question.id},
            citations=[],
            confidence=0.8,
            status=ProposalStatus.PENDING,
        )
    )
    await db_session.flush()

    with pytest.raises(Conflict, match="待审核"):
        await generate_answer(db_session, question.id, "en", actor=actor)


@pytest.mark.asyncio
async def test_scoped_viewer_only_sees_its_engagement(client, db_session):
    lead = await _user(db_session)
    allowed = await _engagement(db_session, lead)
    hidden = AuditEngagement(
        name="Hidden audit",
        audit_type=AuditType.INTERNAL,
        status=EngagementStatus.PREPARING,
        created_by=lead.id,
    )
    db_session.add(hidden)
    await db_session.flush()
    viewer = await _user(db_session, Role.VIEWER)
    viewer.engagement_scope_id = allowed.id
    await db_session.flush()

    visible = await client.get("/api/audit/engagements", headers=_headers(viewer))
    blocked = await client.get(
        f"/api/audit/engagements/{hidden.id}/questions", headers=_headers(viewer)
    )

    assert [row["id"] for row in visible.json()] == [allowed.id]
    assert blocked.status_code == 404


def _question_workbook(rows: list[list[str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["审计问题", "语言"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.mark.asyncio
async def test_excel_import_preserves_each_row_language(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)

    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/questions/xlsx",
        headers=_headers(actor),
        files={
            "file": (
                "questions.xlsx",
                _question_workbook(
                    [
                        ["How is privileged access approved?", "English"],
                        ["如何复核特权账号？", "中文"],
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    assert [(row["question_text"], row["language"]) for row in response.json()] == [
        ("How is privileged access approved?", "en"),
        ("如何复核特权账号？", "zh"),
    ]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_questions.import")
    )
    assert audit.after["source"] == "xlsx"


@pytest.mark.asyncio
async def test_word_export_contains_finalized_answers_only(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    finalized, pending = await import_questions(
        db_session, engagement.id, "Final question?\nPending question?", "en", actor=actor
    )
    answer = AnswerDraft(
        question_id=finalized.id,
        body="Approved response.",
        language="en",
        cited_clause_ids=[],
        cited_control_ids=[],
        suggested_evidence_ids=[],
        gap_notes="",
        reviewed_by=actor.id,
    )
    db_session.add(answer)
    await db_session.flush()
    await finalize_answer(db_session, answer, actor=actor, seen_body=answer.body)

    response = await client.get(
        f"/api/audit/engagements/{engagement.id}/export.docx?language=en",
        headers=_headers(actor),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = WordDocument(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Final question?" in text
    assert "Approved response." in text
    assert pending.question_text not in text
    assert await db_session.scalar(
        select(AuditLog.id).where(AuditLog.action == "audit_engagement.export")
    )


def test_question_similarity_supports_english_and_chinese():
    assert similarity(
        "How is privileged access reviewed?",
        "How do you review privileged access accounts?",
    ) > similarity(
        "How is privileged access reviewed?",
        "How are backups restored?",
    )
    assert similarity("如何复核特权账号？", "特权账号如何定期复核？") > 0.3


@pytest.mark.asyncio
async def test_similar_history_returns_closest_finalized_answer(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    old_question, unrelated, current = await import_questions(
        db_session,
        engagement.id,
        "How is privileged access reviewed?\nHow are backups restored?\nHow do you review privileged access accounts?",
        "en",
        actor=actor,
    )
    for question, body in ((old_question, "Quarterly."), (unrelated, "From snapshots.")):
        answer = AnswerDraft(
            question_id=question.id,
            body=body,
            language="en",
            cited_clause_ids=[],
            cited_control_ids=[],
            suggested_evidence_ids=[],
            gap_notes="",
            reviewed_by=actor.id,
        )
        db_session.add(answer)
        await db_session.flush()
        await finalize_answer(db_session, answer, actor=actor, seen_body=answer.body)

    rows = await similar_history(db_session, current.question_text, exclude_question_id=current.id)
    response = await client.get(
        f"/api/audit/questions/{current.id}/similar-history", headers=_headers(actor)
    )

    assert rows[0].question_id == old_question.id
    assert response.status_code == 200
    assert response.json()[0]["answer"] == "Quarterly."
    assert response.json()[0]["similarity"] > 0


@pytest.mark.asyncio
async def test_batch_generation_enqueues_pending_questions(client, db_session, monkeypatch):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    await import_questions(
        db_session, engagement.id, "Question one?\nQuestion two?", "en", actor=actor
    )
    calls = []

    async def fake_enqueue(function, *args, job_id=None):
        calls.append((function, args, job_id))
        return "audit-job-1"

    monkeypatch.setattr("app.audit.router.enqueue", fake_enqueue)
    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/generate-all",
        headers=_headers(actor),
    )

    assert response.status_code == 200
    assert response.json() == {"job_id": "audit-job-1", "question_count": 2}
    # 任务 id 按项目固定：同一项目还在排队或在跑时，arq 拒绝再投一份（K4）
    assert calls == [("generate_engagement_answers", (engagement.id, actor.id),
                      f"generate-answers:{engagement.id}")]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_answers.batch_enqueue")
    )
    assert audit.after["question_count"] == 2


@pytest.mark.asyncio
async def test_batch_generation_already_running_is_refused_not_duplicated(client, db_session, monkeypatch):
    """双击「批量生成答复」：两个任务同时开始时都看不到对方的待审提案，会各生成一份（K4）。
    同一项目的任务还在排队或在跑时，arq 按固定 id 拒绝入队（返回空）——说清楚，不另投一份。"""
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    await import_questions(db_session, engagement.id, "Question one?", "en", actor=actor)

    async def already_queued(function, *args, job_id=None):
        return ""

    monkeypatch.setattr("app.audit.router.enqueue", already_queued)
    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/generate-all", headers=_headers(actor),
    )
    assert response.status_code == 409
    assert response.json()["code"] == "audit.batch_already_running"
    assert await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_answers.batch_enqueue")
    ) is None


def test_batch_generation_keeps_no_result_so_it_can_run_again():
    """arq 的结果键也会挡住同 id 入队；默认保留一小时，就是一小时内不能再生成一次。"""
    from app.worker import WorkerSettings

    job = next(f for f in WorkerSettings.functions
               if getattr(f, "name", None) == "generate_engagement_answers")
    assert job.keep_result_s == 0


@pytest.mark.asyncio
async def test_preflight_computes_green_yellow_and_red_readiness(client, db_session):
    actor = await _user(db_session)
    framework = Framework(
        key="preview",
        name_zh="预演框架",
        name_en="Preview Framework",
        version="1",
        source="test",
        item_count=3,
        imported_by=actor.id,
    )
    db_session.add(framework)
    await db_session.flush()
    engagement = await _engagement(db_session, actor)
    engagement.framework_id = framework.id
    items = [
        FrameworkItem(
            framework_id=framework.id,
            code=f"PR-{index}",
            title=title,
            description="",
            level=1,
            order_index=index,
        )
        for index, title in enumerate(("Ready", "Implemented", "Gap"), 1)
    ]
    controls = [
        Control(code=f"C-PRE-{index}", title=title, statement=title)
        for index, title in enumerate(("Ready", "Implemented"), 1)
    ]
    db_session.add_all([*items, *controls])
    await db_session.flush()
    db_session.add_all(
        [
            Mapping(
                control_id=controls[index].id,
                framework_item_id=items[index].id,
                strength=MappingStrength.FULL,
                confirmed_by=actor.id,
            )
            for index in range(2)
        ]
    )
    asset = TechAsset(
        name="VaultKeeper",
        category=TechAssetCategory.PAM,
        vendor="VaultKeeper",
        environment=TechAssetEnvironment.PROD,
        status=TechAssetStatus.ACTIVE,
    )
    db_session.add(asset)
    await db_session.flush()
    db_session.add_all(
        [
            Implementation(
                control_id=control.id,
                tech_asset_id=asset.id,
                description="Configured",
                how_enforced=HowEnforced.AUTOMATED,
                status=ImplementationStatus.IMPLEMENTED,
            )
            for control in controls
        ]
    )
    evidence_type = EvidenceType(
        name_zh="报告",
        name_en="Report",
        format="PDF",
        cadence=EvidenceCadence.QUARTERLY,
    )
    db_session.add(evidence_type)
    await db_session.flush()
    db_session.add(
        EvidenceItem(
            evidence_type_id=evidence_type.id,
            control_id=controls[0].id,
            title="Quarterly review",
            last_collected_at=datetime.now(UTC),
            valid_until=datetime.now(UTC) + timedelta(days=30),
            status=EvidenceStatus.COLLECTED,
        )
    )
    await db_session.flush()

    response = await client.get(
        f"/api/audit/engagements/{engagement.id}/preflight?language=en",
        headers=_headers(actor),
    )

    assert response.status_code == 200
    assert response.json()["summary"] == {"green": 1, "yellow": 1, "red": 1}
    assert [row["readiness"] for row in response.json()["rows"]] == [
        "green",
        "yellow",
        "red",
    ]
    assert response.json()["rows"][0]["tool_names"] == ["VaultKeeper"]


@pytest.mark.asyncio
async def test_similar_history_does_not_cross_an_engagement_scope(db_session):
    """限范围账号不得读到别的审计项目的答复。

    `engagement_scope_id` 的定义是「外部审计员账号：限时 + 限某次审计范围」——
    泄露对象是外部人员。`/history` 那个端点按它过滤了，但 `similar_history`
    查的是全部项目的已定稿答复，两个调用点都没传范围：
      · GET /questions/{id}/similar-history 把对方答复全文直接返回；
      · generate_answer 把它喂进提示词，于是对方的内容可能出现在本项目
        交付的答复里，而且没有任何痕迹说明它从哪来。
    """
    from datetime import UTC, datetime

    from app.audit.history import similar_history
    from app.audit.models import QuestionStatus

    actor = await _user(db_session)
    mine = await _engagement(db_session, actor)
    theirs = AuditEngagement(
        name="别人的项目", audit_type=AuditType.EXTERNAL,
        status=EngagementStatus.PREPARING, created_by=actor.id,
    )
    db_session.add(theirs)
    await db_session.flush()

    (asked,) = await import_questions(
        db_session, mine.id, "How is privileged access reviewed?", "en", actor=actor)
    (other,) = await import_questions(
        db_session, theirs.id, "How is privileged access reviewed?", "en", actor=actor)
    other.status = QuestionStatus.FINALIZED
    db_session.add(AnswerDraft(
        question_id=other.id, body="别人项目的机密答复", language="en",
        cited_clause_ids=[], cited_control_ids=[], suggested_evidence_ids=[],
        gap_notes="", reviewed_by=actor.id, finalized_at=datetime.now(UTC),
    ))
    await db_session.flush()

    # 不限范围的内部账号：跨项目复用是这个功能本来的用途，照常能看到
    unscoped = await similar_history(
        db_session, asked.question_text, exclude_question_id=asked.id)
    assert any(h.engagement_name == "别人的项目" for h in unscoped)

    # 限范围账号：一条都不该拿到
    scoped = await similar_history(
        db_session, asked.question_text, exclude_question_id=asked.id,
        scope_engagement_id=mine.id)
    assert scoped == [], f"限范围账号读到了别的项目：{[h.engagement_name for h in scoped]}"


@pytest.mark.asyncio
@pytest.mark.parametrize("rows,code,params", [
    # 同一类问题归到一起报行号，而不是只报第一行
    ([["Q1", "en"], ["", "en"], ["Q3", "en"], ["", "zh"]], "audit.xlsx_questions_blank", {"rows": [3, 5]}),
    ([["Q1", "klingon"]], "audit.xlsx_languages_invalid", {"rows": [2]}),
    ([], "audit.xlsx_no_questions", {}),
])
async def test_excel_import_rejections_carry_a_code(client, db_session, rows, code, params):
    # 原来整段是「第 3 行审计问题为空；第 5 行……」一句中文，套在 audit.import_rejected
    # 的 {{detail}} 里——英文界面只能显示中文。
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/questions/xlsx",
        headers=_headers(actor),
        files={"file": ("questions.xlsx", _question_workbook(rows), "application/octet-stream")},
    )
    assert response.status_code == 400
    assert (response.json()["code"], response.json().get("params") or {}) == (code, params)


@pytest.mark.asyncio
async def test_excel_import_keeps_the_code_of_a_file_it_cannot_read(client, db_session):
    # read_sheet 的拒绝自带 code，原来在路由里被改写成 import_rejected。
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/questions/xlsx",
        headers=_headers(actor),
        files={"file": ("questions.xlsx", b"not excel", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert response.json()["code"] == "matrix.unreadable"


async def _answer(db_session, actor, body="Access is reviewed quarterly."):
    engagement = await _engagement(db_session, actor)
    question = (await import_questions(db_session, engagement.id, "Question?", "en", actor=actor))[0]
    answer = AnswerDraft(
        question_id=question.id, body=body, language="en", cited_clause_ids=[],
        cited_control_ids=[], suggested_evidence_ids=[], gap_notes="", reviewed_by=actor.id,
    )
    db_session.add(answer)
    await db_session.flush()
    return question, answer


@pytest.mark.asyncio
async def test_finalize_freezes_the_text_the_finalizer_saw(client, db_session):
    actor = await _user(db_session)
    _, answer = await _answer(db_session, actor)

    response = await client.post(
        f"/api/audit/answers/{answer.id}/finalize", headers=_headers(actor),
        json={"body": "Access is reviewed quarterly."},
    )

    assert response.status_code == 200
    assert response.json()["final_body"] == "Access is reviewed quarterly."


@pytest.mark.asyncio
async def test_finalize_refuses_when_the_stored_text_is_not_what_was_seen(client, db_session):
    # 定稿原来不带任何内容，后端拿库里的 body 定稿。页面上改了没保存就点定稿，
    # 定下来的（也是导出给审计师的）是旧文本；别人在此期间保存过，定稿人定的是
    # 自己没看过的文本。人确认的是屏幕上那一版，执行的必须是同一版。
    actor = await _user(db_session)
    question, answer = await _answer(db_session, actor)

    response = await client.post(
        f"/api/audit/answers/{answer.id}/finalize", headers=_headers(actor),
        json={"body": "Access is reviewed monthly."},
    )

    assert response.status_code == 409
    assert response.json()["code"] == "audit.response_changed"
    await db_session.refresh(answer)
    await db_session.refresh(question)
    assert answer.finalized_at is None and answer.final_body is None
    assert question.status is not QuestionStatus.FINALIZED


@pytest.mark.asyncio
async def test_finalize_without_the_seen_text_is_rejected(client, db_session):
    actor = await _user(db_session)
    _, answer = await _answer(db_session, actor)
    response = await client.post(f"/api/audit/answers/{answer.id}/finalize", headers=_headers(actor))
    assert response.status_code == 422
    await db_session.refresh(answer)
    assert answer.finalized_at is None


@pytest.mark.asyncio
async def test_saving_an_answer_on_a_stale_base_is_refused(client, db_session):
    # 两人同时编辑同一份答复：后保存的原来直接覆盖先保存的。保存时带上开始编辑时的
    # 那一版，库里已经变了就拒绝。
    actor = await _user(db_session)
    _, answer = await _answer(db_session, actor)

    first = await client.patch(
        f"/api/audit/answers/{answer.id}", headers=_headers(actor),
        json={"body": "Access is reviewed monthly.", "expected_body": "Access is reviewed quarterly."},
    )
    assert first.status_code == 200

    second = await client.patch(
        f"/api/audit/answers/{answer.id}", headers=_headers(actor),
        json={"body": "Access is reviewed weekly.", "expected_body": "Access is reviewed quarterly."},
    )
    assert second.status_code == 409
    assert second.json()["code"] == "audit.response_changed"
    await db_session.refresh(answer)
    assert answer.body == "Access is reviewed monthly."


@pytest.mark.asyncio
async def test_saving_an_answer_without_the_version_it_was_based_on_is_refused(client, db_session):
    """并发保护原来可选：不带 expected_body 的调用方照样能覆盖别人刚存的（K10）。"""
    actor = await _user(db_session)
    _, answer = await _answer(db_session, actor)
    original = answer.body

    response = await client.patch(
        f"/api/audit/answers/{answer.id}", headers=_headers(actor),
        json={"body": "Overwritten blindly."},
    )
    assert response.status_code == 422
    await db_session.refresh(answer)
    assert answer.body == original


@pytest.mark.asyncio
async def test_a_scoped_viewer_cannot_read_other_engagements_answers_through_the_review_queue(
    client, db_session
):
    # 限范围账号（外部审计员）在审计模块里只看得到自己那个项目；但确认队列只要 READ，
    # `/api/proposals?kind=answer` 把**所有**项目待确认的答复草稿——问题 id、正文——
    # 一并交出去。与 similar-history 那次是同一个口子。
    lead = await _user(db_session)
    allowed = await _engagement(db_session, lead)
    hidden = AuditEngagement(name="Hidden audit", audit_type=AuditType.INTERNAL,
                             status=EngagementStatus.PREPARING, created_by=lead.id)
    db_session.add(hidden)
    await db_session.flush()
    mine = (await import_questions(db_session, allowed.id, "Mine?", "en", actor=lead))[0]
    theirs = (await import_questions(db_session, hidden.id, "Theirs?", "en", actor=lead))[0]
    for question, body in ((mine, "Visible draft."), (theirs, "Confidential draft.")):
        db_session.add(Proposal(
            kind=ProposalKind.ANSWER, status=ProposalStatus.PENDING, confidence=0.5, citations=[],
            payload={"question_id": question.id, "body": body},
        ))
    viewer = await _user(db_session, Role.VIEWER)
    viewer.engagement_scope_id = allowed.id
    await db_session.flush()

    rows = (await client.get("/api/proposals?kind=answer", headers=_headers(viewer))).json()
    assert [row["payload"]["question_id"] for row in rows] == [mine.id]

    everything = (await client.get("/api/proposals", headers=_headers(viewer))).json()
    assert "Confidential draft." not in str(everything)

    # 不限范围的账号照旧看得到全部
    unscoped = (await client.get("/api/proposals?kind=answer", headers=_headers(lead))).json()
    assert {row["payload"]["question_id"] for row in unscoped} == {mine.id, theirs.id}
