"""有 Heading 样式的 Word 文档里，正文样式手打的下级编号标题（KNOWN_ISSUES P1）。

真实文档常见这种写法：`6` 与 `6.1` 用 Heading 样式（编号由 numbering.xml 自动画），
`6.1.1 Access Requests` 却是在正文样式里手打的。原来整节塌成一条，引用给不到 `6.1.1`。

以下仅使用合成文档验证编号机制，不代表真实语料验收（KNOWN_ISSUES P1）。
"""

from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml

from app.parsing.docx_parser import DocxParser

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_LEVELS = "".join(
    f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    f'<w:lvlText w:val="{".".join("%" + str(k + 1) for k in range(i + 1))}"/></w:lvl>'
    for i in range(4)
)
_NUMBERING = (
    f'<w:numbering {_W}><w:abstractNum w:abstractNumId="0">{_LEVELS}</w:abstractNum>'
    '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>'
)


def _policy(tmp_path: Path, extra: list[str] = (), *, auto: bool = True) -> Path:
    """第 6 章：`6` / `6.1` / `6.2` 用 Heading 样式，`6.1.x` 手打在正文样式里。

    auto=True 时标题编号由 numbering.xml 渲染（文字里没有编号），
    auto=False 时编号直接写在标题文字里——两种写法都常见。
    """
    doc = Document()
    if auto:
        # python-docx 的默认模板自带 numbering.xml，另挂一个会被忽略——得替换掉原有的。
        try:
            doc.part.numbering_part._element = parse_xml(_NUMBERING)
        except NotImplementedError:
            doc.part.relate_to(Part(
                PackURI("/word/numbering.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
                _NUMBERING.encode(), doc.part.package,
            ), RT.NUMBERING)

    def heading(number: str, title: str, level: int) -> None:
        paragraph = doc.add_heading(title if auto else f"{number} {title}", level=level)
        if auto:
            paragraph._p.get_or_add_pPr().append(parse_xml(
                f'<w:numPr {_W}><w:ilvl w:val="{level - 1}"/><w:numId w:val="1"/></w:numPr>'
            ))

    for chapter in range(1, 6):
        heading(str(chapter), f"Topic {chapter}", 1)
        doc.add_paragraph(f"Teams must follow topic {chapter}.")
    heading("6", "Access Control", 1)
    doc.add_paragraph("This section sets access rules.")
    heading("6.1", "User Access", 2)
    doc.add_paragraph("Access is granted on least privilege.")
    for index, title in enumerate(["Access Requests", "Access Reviews", "Privileged Access"], 1):
        doc.add_paragraph(f"6.1.{index} {title}")
        doc.add_paragraph(f"Owners must handle {title.lower()} and keep evidence.")
    for line in extra:
        doc.add_paragraph(line)
    heading("6.2", "Remote Access", 2)
    doc.add_paragraph("Remote access must use MFA.")
    path = tmp_path / "policy.docx"
    doc.save(path)
    return path


def _walk(nodes):
    for node in nodes:
        yield node
        yield from _walk(node.children)


def _chapter(parsed, number: str):
    return next(node for node in parsed.clauses if node.number == number)


@pytest.mark.parametrize("auto", [True, False], ids=["auto-numbered", "typed-number"])
def test_typed_subheadings_hang_under_the_real_styled_parent(tmp_path, auto):
    parsed = DocxParser().parse(_policy(tmp_path, auto=auto))
    chapter = _chapter(parsed, "6")
    assert [(n.number, n.heading) for n in chapter.children] == [
        ("6.1", "User Access"), ("6.2", "Remote Access"),
    ]
    user_access = chapter.children[0]
    # 父节点保住自己的正文：它是子条款进抽取时的上下文
    assert user_access.text == "Access is granted on least privilege."
    assert [(n.number, n.heading, n.text) for n in user_access.children] == [
        ("6.1.1", "Access Requests", "Owners must handle access requests and keep evidence."),
        ("6.1.2", "Access Reviews", "Owners must handle access reviews and keep evidence."),
        ("6.1.3", "Privileged Access", "Owners must handle privileged access and keep evidence."),
    ]
    numbers = [n.number for n in _walk(parsed.clauses) if n.number]
    assert len(numbers) == len(set(numbers)), "不能为每个子标题各补一个空的 6.1"


def test_three_level_numbering_without_styles_does_not_duplicate_the_parent(tmp_path):
    """同一个补祖先的循环也管没有样式的文档：`4.1.1` 前面明明就是 `4.1`。"""
    doc = Document()
    for line in ["4 Access Control", "Intro.", "4.1 User Access", "Least privilege.",
                 "4.1.1 Requests", "Body a.", "4.1.2 Reviews", "Body b.", "4.2 Remote", "MFA."]:
        doc.add_paragraph(line)
    path = tmp_path / "plain.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)
    assert [n.number for n in parsed.clauses] == ["4"]
    user_access = parsed.clauses[0].children[0]
    assert (user_access.number, user_access.text) == ("4.1", "Least privilege.")
    assert [n.number for n in user_access.children] == ["4.1.1", "4.1.2"]
    assert [n.number for n in parsed.clauses[0].children] == ["4.1", "4.2"]


def test_a_numbered_cross_reference_stays_in_the_body(tmp_path):
    """`2.1` 接不上当前的 `6.1.x`：它是引用，不是标题。当成标题会在第 6 章中间补出一个第 2 章。"""
    parsed = DocxParser().parse(_policy(tmp_path, [
        "2.1 of the Standard applies to vendors.",
    ]))
    assert [n.number for n in parsed.clauses if n.number] == ["1", "2", "3", "4", "5", "6"]
    last = _chapter(parsed, "6").children[0].children[-1]
    assert last.number == "6.1.3"
    assert "2.1 of the Standard applies to vendors." in last.text


def test_a_sentence_opening_with_section_n_stays_in_the_body(tmp_path):
    parsed = DocxParser().parse(_policy(tmp_path, [
        "Section 5 applies to all vendors.",
    ]))
    last = _chapter(parsed, "6").children[0].children[-1]
    assert "Section 5 applies to all vendors." in last.text
    assert [n.number for n in parsed.clauses if n.number] == ["1", "2", "3", "4", "5", "6"]


def test_a_number_missing_its_chapter_is_warned_about_not_guessed(tmp_path):
    parsed = DocxParser().parse(_policy(tmp_path, [
        ".6.1.4 Service Accounts", "Service accounts need an owner.",
    ]))
    last = _chapter(parsed, "6").children[0].children[-1]
    assert last.number == "6.1.3"
    assert ".6.1.4 Service Accounts" in last.text
    warning = next(w for w in parsed.warnings if getattr(w, "code", None) == "parse.suspect_numbering")
    assert warning.params == {"count": 1}


def test_a_clean_document_raises_no_numbering_warning(tmp_path):
    parsed = DocxParser().parse(_policy(tmp_path))
    assert not [w for w in parsed.warnings if getattr(w, "code", None) == "parse.suspect_numbering"]
