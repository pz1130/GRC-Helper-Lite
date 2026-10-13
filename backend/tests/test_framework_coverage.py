import pytest

from app.controls.models import Control
from app.frameworks.coverage import gaps, is_requirement, summarize
from app.frameworks.models import Framework, FrameworkItem, Mapping, MappingStrength
from app.review.models import Proposal, ProposalKind, ProposalStatus


def _item(**kwargs):
    from types import SimpleNamespace
    return SimpleNamespace(**{"attributes": None, **kwargs})


def test_a_leaf_is_a_requirement_by_default():
    assert is_requirement(_item(), has_children=False) is True


def test_a_container_is_not_a_requirement_by_default():
    assert is_requirement(_item(), has_children=True) is False


def test_an_explicit_flag_overrides_the_leaf_rule_in_both_directions():
    assert is_requirement(_item(attributes={"is_requirement": True}), has_children=True) is True
    assert is_requirement(_item(attributes={"is_requirement": False}), has_children=False) is False


async def _corpus(db_session):
    fw = Framework(key="k", name_zh="z", name_en="e", version="1", source="s", item_count=0)
    control = Control(code="C-0001", title="t", statement="s")
    db_session.add_all([fw, control])
    await db_session.flush()

    root = FrameworkItem(framework_id=fw.id, code="AC", title="Access Control",
                         description="", level=1, order_index=0)
    db_session.add(root)
    await db_session.flush()
    covered = FrameworkItem(
        framework_id=fw.id, parent_id=root.id, code="AC-1", title="Policy",
        description="d", level=2, order_index=1,
        attributes={"baselines": ["low", "moderate"]})
    supported_only = FrameworkItem(
        framework_id=fw.id, parent_id=root.id, code="AC-2", title="Accounts",
        description="d", level=2, order_index=2,
        attributes={"baselines": ["moderate"]})
    untouched = FrameworkItem(
        framework_id=fw.id, parent_id=root.id, code="AC-3", title="Enforcement",
        description="d", level=2, order_index=3,
        attributes={"baselines": ["high"]})
    db_session.add_all([covered, supported_only, untouched])
    await db_session.flush()

    db_session.add(Mapping(control_id=control.id, framework_item_id=covered.id,
                           strength=MappingStrength.PARTIAL, rationale="r", quote="d"))
    db_session.add(Mapping(control_id=control.id, framework_item_id=supported_only.id,
                           strength=MappingStrength.SUPPORTING, rationale="r", quote="d"))
    await db_session.flush()
    return fw, control, covered, supported_only, untouched


@pytest.mark.asyncio
async def test_supporting_alone_does_not_close_a_gap(db_session):
    fw, _, _, _supported_only, _untouched = await _corpus(db_session)
    rows = await gaps(db_session, fw.id)
    codes = {row.code for row in rows}
    assert codes == {"AC-2", "AC-3"}
    assert next(row for row in rows if row.code == "AC-2").has_supporting is True
    assert next(row for row in rows if row.code == "AC-3").has_supporting is False


@pytest.mark.asyncio
async def test_containers_are_excluded_from_the_denominator(db_session):
    fw, *_ = await _corpus(db_session)
    rows = await summarize(db_session, fw.id)
    root = next(row for row in rows if row.code == "AC")
    assert root.requirements == 3
    assert root.covered == 1


@pytest.mark.asyncio
async def test_baseline_filters_the_denominator(db_session):
    fw, *_ = await _corpus(db_session)
    rows = await summarize(db_session, fw.id, baseline="moderate")
    root = next(row for row in rows if row.code == "AC")
    assert root.requirements == 2
    assert root.covered == 1
    assert {row.code for row in await gaps(db_session, fw.id, baseline="moderate")} == {"AC-2"}


@pytest.mark.asyncio
async def test_pending_proposals_never_affect_coverage(db_session):
    fw, control, _, _, untouched = await _corpus(db_session)
    db_session.add(Proposal(
        kind=ProposalKind.MAPPING, status=ProposalStatus.PENDING,
        payload={"framework_item_id": untouched.id, "control_id": control.id,
                 "strength": "full", "quote": "d", "rationale": "r", "confidence": 0.99},
        citations=[], confidence=0.99,
    ))
    await db_session.flush()
    assert "AC-3" in {row.code for row in await gaps(db_session, fw.id)}


@pytest.mark.asyncio
async def test_an_empty_framework_has_no_requirements(db_session):
    fw = Framework(key="empty", name_zh="z", name_en="e", version="1", source="s", item_count=0)
    db_session.add(fw)
    await db_session.flush()
    assert await summarize(db_session, fw.id) == []
    assert await gaps(db_session, fw.id) == []


async def _superseded_corpus(db_session):
    """R1 只由出处全在已替换文档里的控制点 A 关上；其余三项都还有依据。"""
    from app.clauses.models import Clause
    from app.controls.models import ControlSource, SourceRelation
    from app.ingest.models import DocStatus, DocType, Document

    fw = Framework(key="k", name_zh="", name_en="e", version="1", source="s", item_count=0)
    old = Document(title="Policy v1", doc_type=DocType.POLICY, file_hash="1" * 64, file_path="/1",
                   original_filename="1.pdf", status=DocStatus.SUPERSEDED)
    new = Document(title="Policy v2", doc_type=DocType.POLICY, file_hash="2" * 64, file_path="/2",
                   original_filename="2.pdf", status=DocStatus.ACTIVE)
    a, b, c = (Control(code=code, title=code, statement="s") for code in ("C-A", "C-B", "C-C"))
    db_session.add_all([fw, old, new, a, b, c])
    await db_session.flush()

    def clause(doc, number):
        return Clause(document_id=doc.id, number=number, heading="H", heading_path=f"H › {number}",
                      citation_label=number, text="t", order_index=0, level=1)

    old_clause, other_old, new_clause = clause(old, "1"), clause(old, "2"), clause(new, "2")
    db_session.add_all([old_clause, other_old, new_clause])
    await db_session.flush()
    db_session.add_all([
        ControlSource(control_id=a.id, clause_id=old_clause.id, relation=SourceRelation.DEFINES),
        ControlSource(control_id=b.id, clause_id=other_old.id, relation=SourceRelation.DEFINES),
        ControlSource(control_id=b.id, clause_id=new_clause.id, relation=SourceRelation.DEFINES),
    ])

    root = FrameworkItem(framework_id=fw.id, code="R", title="Root", description="", level=1, order_index=0)
    db_session.add(root)
    await db_session.flush()
    items = {}
    for index, code in enumerate(("R1", "R2", "R3", "R4"), start=1):
        items[code] = FrameworkItem(framework_id=fw.id, parent_id=root.id, code=code, title=code,
                                    description="", level=2, order_index=index)
    db_session.add_all(items.values())
    await db_session.flush()
    for control, code in ((a, "R1"), (b, "R2"), (c, "R3"), (a, "R4"), (c, "R4")):
        db_session.add(Mapping(control_id=control.id, framework_item_id=items[code].id,
                               strength=MappingStrength.FULL, rationale="r", quote="q"))
    await db_session.flush()
    return fw


@pytest.mark.asyncio
async def test_a_requirement_backed_only_by_superseded_policies_is_flagged_not_uncovered(db_session):
    # 新版删掉了某条要求、旧版被替换后，依据它的控制点与映射都还在，覆盖度照样是绿的。
    # 定下的做法是：仍算覆盖（不替人改已确认的数据），但单独标出来。
    from app.frameworks.coverage import superseded_support

    fw = await _superseded_corpus(db_session)

    root = next(row for row in await summarize(db_session, fw.id) if row.code == "R")
    assert (root.requirements, root.covered, root.superseded_only) == (4, 4, 1)
    assert [(row.code, row.controls) for row in await superseded_support(db_session, fw.id)] == [
        ("R1", ["C-A"]),
    ]
    assert await gaps(db_session, fw.id) == []


@pytest.mark.asyncio
async def test_superseded_support_is_served_to_the_framework_page(client, db_session):
    from app.iam.models import User
    from app.iam.permissions import Role
    from app.iam.security import hash_password

    fw = await _superseded_corpus(db_session)
    db_session.add(User(email="v@example.com", name="v", role=Role.VIEWER, password_hash=hash_password("pw123456")))
    await db_session.flush()
    token = (await client.post("/api/auth/login", json={"email": "v@example.com", "password": "pw123456"})).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    rows = (await client.get(f"/api/frameworks/{fw.id}/superseded-support", headers=headers)).json()
    assert [(row["code"], row["controls"]) for row in rows] == [("R1", ["C-A"])]
    coverage = (await client.get(f"/api/frameworks/{fw.id}/coverage", headers=headers)).json()
    assert next(row for row in coverage if row["code"] == "R")["superseded_only"] == 1


@pytest.mark.asyncio
async def test_the_readiness_package_carries_the_superseded_flag(db_session):
    # 只在页面上标「依据已作废」，导出的准备包里覆盖度照样是满的，看不出来。
    import io
    import zipfile

    from docx import Document as WordDocument

    from app.packaging.service import build_package

    fw = await _superseded_corpus(db_session)
    content = await build_package(db_session, fw.id, language="en")
    with zipfile.ZipFile(io.BytesIO(content)) as bundle:
        word = WordDocument(io.BytesIO(bundle.read("readiness.docx")))
    text = "\n".join(p.text for p in word.paragraphs)
    cells = [cell.text for table in word.tables for row in table.rows for cell in row.cells]

    assert "Covered only by superseded policies" in text
    assert "Of which on superseded policies only" in cells
    assert "R1" in cells and "C-A" in cells
