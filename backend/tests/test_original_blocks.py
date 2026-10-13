"""原文快照的保留、位置和拆合传递，使用合成文本。"""

import pytest
from docx import Document as WordDocument
from sqlalchemy import func, select

from app.clauses.edit import merge_clauses, split_clause
from app.clauses.models import Clause, ClauseBlockSource, OriginalBlock
from app.clauses.sources import source_page
from app.parsing.contract import ClauseNode, SourceBlock
from app.parsing.docx_parser import DocxParser
from app.parsing.flatten import persist
from app.parsing.pdf_parser import PdfParser
from tests.test_clause_edit import _doc, _user
from tests.test_parsing_pdf import _write_pdf


def tree(text="Review access. Keep records."):
    return [ClauseNode(heading="Review", text=text, level=1, source_keys=["p1"])]


def blocks(text="Review access. Keep records."):
    return [SourceBlock(key="p1", text=text, kind="paragraph", locator={"body_index": 1})]


async def save(session, doc, nodes=None, originals=None):
    await persist(session, tree() if nodes is None else nodes, document_id=doc.id,
                  source_blocks=blocks() if originals is None else originals,
                  source_origin="docx_xml")


def test_word_sources_preserve_spacing_and_merged_cell_locations(tmp_path):
    document = WordDocument()
    document.add_paragraph("  Preserve original spacing.  ")
    document.add_heading("Review", 1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shared heading"
    table.cell(1, 0).text = "Frequency"
    table.cell(1, 1).text = "Quarterly"
    path = tmp_path / "merged.docx"
    document.save(path)
    parsed = DocxParser().parse(path)
    assert parsed.source_origin == "docx_xml"
    assert parsed.source_blocks[0].text == "  Preserve original spacing.  "
    cells = [block for block in parsed.source_blocks if block.kind == "table_cell"]
    assert len(cells) == 3
    assert cells[0].locator["positions"] == [{"row": 1, "column": 1}, {"row": 1, "column": 2}]
    table_node = parsed.clauses[1].children[0]
    assert set(table_node.source_keys) == {block.key for block in cells}


def test_pdf_sources_keep_real_pages_and_link_to_clauses(tmp_path):
    parsed = PdfParser().parse(_write_pdf(tmp_path / "source.pdf", [
        ["Before the heading", "1 Review", "Review access."], ["2 Records", "Keep records."],
    ]))
    assert parsed.source_origin == "pdf_text"
    assert [block.locator["page"] for block in parsed.source_blocks] == [1, 1, 1, 2, 2]
    assert {key for node in parsed.clauses for key in node.source_keys} == {
        block.key for block in parsed.source_blocks
    }


async def test_reparse_reuses_originals_and_changed_input_keeps_old_snapshot(db_session):
    document = await _doc(db_session)
    await save(db_session, document)
    old = await db_session.scalar(select(OriginalBlock))
    old_id, old_text = old.id, old.text
    await save(db_session, document)
    assert await db_session.scalar(select(func.count()).select_from(OriginalBlock)) == 1
    assert (await source_page(db_session, document.id))["blocks"][0]["id"] == old_id
    await save(db_session, document, tree("Review weekly."), blocks("Review weekly."))
    await db_session.refresh(old)
    assert old.text == old_text
    assert await db_session.scalar(select(func.count()).select_from(OriginalBlock)) == 2
    page = await source_page(db_session, document.id)
    assert page["total"] == 1
    assert page["blocks"][0]["text"] == "Review weekly."


async def test_split_and_merge_preserve_original_text_and_block_ids(db_session):
    document, actor = await _doc(db_session), await _user(db_session)
    await save(db_session, document)
    clause = await db_session.scalar(select(Clause))
    before = (await source_page(db_session, document.id, clause_id=clause.id))["blocks"]
    right = await split_clause(db_session, clause.id, at=clause.text.index("Keep"), actor=actor)
    for cid in (clause.id, right.id):
        page = await source_page(db_session, document.id, clause_id=cid)
        assert page["blocks"][0]["id"] == before[0]["id"]
        assert page["blocks"][0]["text"] == before[0]["text"]
    await merge_clauses(db_session, winner_id=clause.id, loser_id=right.id, actor=actor)
    page = await source_page(db_session, document.id, clause_id=clause.id)
    assert page["total"] == 1
    assert page["blocks"][0]["clause_ids"] == [clause.id]
    assert await db_session.scalar(select(func.count()).select_from(ClauseBlockSource)) == 2


async def test_unassigned_blocks_are_retained_and_pagination_is_scoped(db_session):
    document = await _doc(db_session)
    originals = [*blocks(), SourceBlock("p2", "Unassigned requirement.", "paragraph", {})]
    await save(db_session, document, originals=originals)
    page = await source_page(db_session, document.id, limit=1)
    assert page["total"] == 2 and page["unassigned"] == 1 and page["has_more"]
    missing = await source_page(db_session, document.id, unassigned_only=True)
    assert [block["text"] for block in missing["blocks"]] == ["Unassigned requirement."]
    assert not missing["has_more"]
    assert (await source_page(db_session, document.id, offset=2))["blocks"] == []
    other = await _doc(db_session, file_hash="b" * 64)
    assert (await source_page(db_session, other.id))["available"] is False
    assert (await source_page(db_session, other.id))["blocks"] == []
    await save(db_session, other, [], [])
    assert (await source_page(db_session, other.id))["available"] is True
    assert (await source_page(db_session, other.id))["total"] == 0


async def test_invalid_provenance_rolls_back_with_the_parse(db_session):
    document = await _doc(db_session)
    await save(db_session, document)
    original_snapshot = document.source_snapshot
    with pytest.raises(KeyError):
        async with db_session.begin_nested():
            await save(db_session, document,
                       [ClauseNode("Broken", "Changed text", 1, source_keys=["absent"])],
                       blocks("Changed text"))
    await db_session.refresh(document)
    assert document.source_snapshot == original_snapshot
    assert await db_session.scalar(select(func.count()).select_from(OriginalBlock)) == 1
    assert (await db_session.scalar(select(Clause))).text == "Review access. Keep records."


async def test_source_endpoints_require_login_and_allow_readers(client, db_session):
    from app.iam.permissions import Role

    document = await _doc(db_session)
    await save(db_session, document)
    clause = await db_session.scalar(select(Clause))
    endpoint = f"/api/documents/{document.id}/source-blocks"
    assert (await client.get(endpoint)).status_code == 401
    user = await _user(db_session, Role.VIEWER)
    login = await client.post("/api/auth/login", json={"email": user.email, "password": "pw123456"})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    for path in (endpoint, f"/api/clauses/{clause.id}/source-blocks"):
        response = await client.get(path, headers=headers)
        assert response.status_code == 200
        assert response.json()["blocks"][0]["text"] == "Review access. Keep records."
        assert (await client.get(path + "?limit=201", headers=headers)).status_code == 422
    assert (await client.get("/api/clauses/999999/source-blocks", headers=headers)).status_code == 404


async def test_duplicate_keys_are_rejected_before_replacing_clauses(db_session):
    document = await _doc(db_session)
    await save(db_session, document)
    with pytest.raises(ValueError, match="unique"):
        await save(db_session, document, originals=blocks() + blocks())
    assert await db_session.scalar(select(func.count()).select_from(Clause)) == 1


async def test_plain_text_fallback_is_a_separate_source_snapshot(client, db_session, monkeypatch):
    from unittest.mock import AsyncMock

    document = await _doc(db_session)
    await save(db_session, document)
    old_block = await db_session.scalar(select(OriginalBlock))
    user = await _user(db_session)
    login = await client.post("/api/auth/login", json={"email": user.email, "password": "pw123456"})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    monkeypatch.setattr("app.ingest.router.enqueue", AsyncMock(return_value="queued"))
    response = await client.post(f"/api/documents/{document.id}/plain-text",
                                 json={"text": "Owners must review monthly.\nKeep records."},
                                 headers=headers)
    assert response.status_code == 200
    page = await source_page(db_session, document.id)
    assert page["total"] == 2 and page["unassigned"] == 0
    assert {block["origin"] for block in page["blocks"]} == {"pasted_text"}
    assert [block["locator"]["line"] for block in page["blocks"]] == [1, 2]
    await db_session.refresh(old_block)
    assert old_block.text == "Review access. Keep records."
