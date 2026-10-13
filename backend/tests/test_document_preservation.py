"""合成文档验证原文保留，不依赖某家机构的样本或真实模型。"""

from docx import Document

from app.parsing.docx_parser import DocxParser
from app.parsing.numbering import assemble_tree
from app.parsing.pdf_parser import PdfParser
from tests.test_parsing_pdf import _write_pdf


def walk(nodes):
    for node in nodes:
        yield node
        yield from walk(node.children)


def test_short_pdf_with_text_is_not_discarded_as_a_scan(tmp_path):
    path = _write_pdf(tmp_path / "short.pdf", [["Owners must report incidents within two days."]])
    parsed = PdfParser().parse(path)
    assert parsed.clauses[0].text == "Owners must report incidents within two days."
    assert parsed.clauses[0].number is None
    assert parsed.clauses[0].page_ref == 1
    assert "parse.scanned_no_text_layer" not in [getattr(w, "code", None) for w in parsed.warnings]


def test_no_heading_pdf_preserves_every_page_in_order(tmp_path):
    pages = [["Owners must review access quarterly.", "Retain review records for three years."],
             ["Report exceptions within two days."]]
    parsed = PdfParser().parse(_write_pdf(tmp_path / "flat.pdf", pages))
    assert [node.text for node in parsed.clauses] == ["\n".join(page) for page in pages]
    assert [node.page_ref for node in parsed.clauses] == [1, 2]
    assert all(node.kind == "unstructured" for node in parsed.clauses)


def test_unstructured_long_text_is_bounded_without_truncation():
    lines = [f"Paragraph {i}: " + "x" * 1000 for i in range(12)]
    nodes = assemble_tree([], lines)
    assert len(nodes) > 1
    assert "\n".join(node.text for node in nodes) == "\n".join(lines)
    assert all(len(node.text) <= 4000 for node in nodes)


def test_word_preamble_and_text_around_tables_are_preserved(tmp_path):
    doc = Document()
    doc.add_paragraph("Owners must keep records.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Three years"
    doc.add_paragraph("Exceptions require approval.")
    doc.add_heading("Reporting", 1)
    doc.add_paragraph("Report annually.")
    path = tmp_path / "preamble.docx"
    doc.save(path)
    nodes = DocxParser().parse(path).clauses
    assert [node.kind for node in nodes] == ["unstructured", "table", "unstructured", "section"]
    assert [node.text for node in nodes] == [
        "Owners must keep records.", "Three years", "Exceptions require approval.", "Report annually.",
    ]


def test_word_mixed_styled_and_normal_numbered_headings(tmp_path):
    doc = Document()
    doc.add_heading("1 Access", 1)
    doc.add_paragraph("Applies to production accounts.")
    doc.add_paragraph("1.1 Review")
    doc.add_paragraph("Owners must review access quarterly.")
    doc.add_heading("2 Records", 1)
    path = tmp_path / "mixed.docx"
    doc.save(path)
    nodes = DocxParser().parse(path).clauses
    assert [(node.number, node.heading) for node in nodes] == [("1", "Access"), ("2", "Records")]
    assert nodes[0].text == "Applies to production accounts."
    assert [(node.number, node.text) for node in nodes[0].children] == [
        ("1.1", "Owners must review access quarterly."),
    ]


def test_word_unnumbered_style_can_parent_numbered_subsections(tmp_path):
    doc = Document()
    doc.add_heading("Access", 1)
    doc.add_paragraph("1.1 Review")
    doc.add_paragraph("Review quarterly.")
    doc.add_paragraph("1.2 Records")
    doc.add_paragraph("Keep records.")
    path = tmp_path / "unnumbered.docx"
    doc.save(path)
    nodes = DocxParser().parse(path).clauses
    assert len(nodes) == 1
    assert [node.number for node in nodes[0].children] == ["1.1", "1.2"]


def test_word_styled_document_keeps_operational_lists_in_body(tmp_path):
    doc = Document()
    doc.add_heading("Steps", 1)
    steps = ["1. Open the form", "2. Fill in the owner", "3. Save the record"]
    for step in steps:
        doc.add_paragraph(step)
    path = tmp_path / "steps.docx"
    doc.save(path)
    nodes = DocxParser().parse(path).clauses
    assert len(list(walk(nodes))) == 1
    assert nodes[0].text == "\n".join(steps)


def _cover_then_sections(tmp_path, cover: list[str], sections: int):
    doc = Document()
    for line in cover:
        doc.add_paragraph(line)
    for index in range(1, sections + 1):
        doc.add_heading(f"Topic {index}", 1)
        doc.add_paragraph(f"Teams must follow topic {index} and keep the evidence for audit.")
    path = tmp_path / "cover.docx"
    doc.save(path)
    return DocxParser().parse(path)


def test_a_short_cover_page_raises_no_structure_warning(tmp_path):
    """封面、缩略语、分发名单不属于任何章节是对的——为它告警等于每份制度都告警。"""
    parsed = _cover_then_sections(
        tmp_path, ["INTERNAL", "Acme Access Procedure", "Version: 1.0", "Distribution List"], 6,
    )
    assert parsed.clauses[0].kind == "unstructured"
    assert "parse.unstructured_text_preserved" not in [w.code for w in parsed.warnings]


def test_a_long_run_of_text_before_the_first_heading_is_still_warned_about(tmp_path):
    """开头几章的标题没认出来时，正文整段落在第一个标题之前——这正是要告警的。"""
    body = [f"Owners must review access for system {i} every quarter." for i in range(30)]
    parsed = _cover_then_sections(tmp_path, body, 2)
    assert "parse.unstructured_text_preserved" in [w.code for w in parsed.warnings]
