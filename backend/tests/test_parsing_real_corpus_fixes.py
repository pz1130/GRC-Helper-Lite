"""真实制度语料（evals/real）暴露问题的回归测试。"""
import importlib.util
import sys
from pathlib import Path

from app.parsing.contract import ClauseNode
from app.parsing.pdf_parser import _zero_point_chapters
from app.parsing.restructure import nest_chapters

_spec = importlib.util.spec_from_file_location(
    "seg_score", Path(__file__).resolve().parents[1] / "evals" / "segmentation" / "score.py")
seg = importlib.util.module_from_spec(_spec)
sys.modules["seg_score"] = seg
_spec.loader.exec_module(seg)


def _n(text, number=None, heading="h"):
    return ClauseNode(heading=heading, text=text, level=1, number=number)


def test_zero_point_chapters_become_top_level_numbers():
    out = _zero_point_chapters(["1.0 Purpose", "body", "2.0 Scope", "3.0 Policy"])
    assert out[0] == "1. Purpose" and out[2] == "2. Scope" and out[1] == "body"


def test_single_zero_point_line_is_left_alone():
    assert _zero_point_chapters(["1.0 Purpose", "body"]) == ["1.0 Purpose", "body"]


def test_chapter_references_inside_sentences_do_not_nest():
    # 章标记落在上一节点正文末尾；句中引用（「Chapter VII.」「Chapter IX (」）不能当章。
    nodes = [_n("intro\nCHAPTER I\nGeneral"), _n("as laid down in Chapter VII.", "1"),
             _n("x\nCHAPTER II\nPrinciples", "2"), _n("see Chapter IX (cooperation)", "3")]
    out = nest_chapters(nodes)
    chapters = [o for o in out if o.number and o.number.startswith("Chapter")]
    assert [c.number for c in chapters] == ["Chapter I", "Chapter II"]
    assert [c.number for c in chapters[0].children] == ["1", "2"]
    assert [c.number for c in chapters[1].children] == ["3"]


def test_part_then_chapter_nesting():
    nodes = [_n("intro\nPART 1\nGeneral"), _n("x", "1"), _n("y\nPART 2\nDue diligence\nCHAPTER 1\nGeneral", "2"),
             _n("z\nCHAPTER 2\nEnhanced", "3"), _n("w", "4")]
    out = nest_chapters(nodes)
    parts = [o for o in out if o.number and o.number.startswith("Part")]
    assert [p.number for p in parts] == ["Part 1", "Part 2"]
    assert [c.number for c in parts[1].children] == ["Chapter 1", "Chapter 2"]


def test_scorer_matches_numbered_paragraph_by_number():
    gold = {"heading": "Employees must report incidents", "number": "4.1.1"}
    assert seg._heading_matches(gold, ClauseNode(heading="4.1.1", text="", level=3, number="4.1.1"))
    assert not seg._heading_matches(gold, ClauseNode(heading="4.1.2", text="", level=3, number="4.1.2"))


def test_gold_nests_articles_under_sections():
    from evals.real.gold import nest_articles_in_sections
    labels = [{"number": "第三章", "level": 1}, {"number": "第一节", "level": 2},
              {"number": "第十二条", "level": 2}, {"number": "第四章", "level": 1}, {"number": "第十三条", "level": 2}]
    assert [x["level"] for x in nest_articles_in_sections(labels)] == [1, 2, 3, 1, 2]


def test_toc_line_keeps_numbered_noun_heading():
    from app.parsing.numbering import is_toc_line
    assert not is_toc_line("3 Appendix 2")
    assert is_toc_line("4.1 Access control 12")


def test_uk_schedules_become_their_own_subtree():
    from app.parsing.numbering import NumberedHeading
    from app.parsing.pdf_parser import _schedules
    lines = ["2 Body", "text", "SCHEDULE", "Regulation 2(1)", "EXEMPT PROCESSING", "1 Interpretation", "x"]
    paragraphs = [NumberedHeading("1", "Interpretation", 5, 6)]
    out = _schedules(lines, 1, paragraphs)
    assert [(h.number, h.title) for h in out] == [("Sch1", "EXEMPT PROCESSING"), ("Sch1.1", "Interpretation")]


def test_heading_styled_long_sentence_is_body(tmp_path):
    import docx

    from app.parsing.docx_parser import _heading_level
    document = docx.Document()
    short = document.add_paragraph("What you can expect from us", style="Heading 2")
    long = document.add_paragraph(
        "When you choose to share your contact information with us, we commit to coordinating "
        "with you as openly and as quickly as possible.", style="Heading 2")
    assert _heading_level(short) == 2
    assert _heading_level(long) is None


def test_bold_refinement_drops_plain_numbered_body_and_adds_front_matter():
    from app.parsing.numbering import NumberedHeading
    from app.parsing.pdf_parser import _refine_by_bold_headings
    b, r = ["X-Bold"], ["X-Regular"]
    rows = [("Abstract", b), ("text", r), ("Keywords", b), ("text", r),
            ("1. Intro", b), ("text", r), ("2.0 website and more words", r),
            ("2. Scope", b), ("text", r), ("3. Use", b), ("text", r), ("4. More", b), ("text", r), ("5. End", b), ("text", r),
            ("Appendix A. Terms", b), ("t", r)]
    lines = [t for t, _ in rows]
    layout = [{"font_size": 11.0, "fonts": f, "bbox": (72, 0, 0, 0)} for _, f in rows]
    heads = [NumberedHeading("1", "Intro", 4), NumberedHeading("2", "website", 6),
             NumberedHeading("2", "Scope", 7), NumberedHeading("3", "Use", 9),
             NumberedHeading("4", "More", 11), NumberedHeading("5", "End", 13)]
    out = _refine_by_bold_headings(heads, lines, layout)
    assert [h.title for h in out] == ["Abstract", "Keywords", "Intro", "Scope", "Use", "More", "End", "Terms"]


def test_single_bold_title_is_not_front_matter():
    from app.parsing.numbering import NumberedHeading
    from app.parsing.pdf_parser import _refine_by_bold_headings
    b, r = ["X-Bold"], ["X-Regular"]
    rows = [("Policy", b), ("text", r), ("1. A", b), ("t", r), ("2. B", b), ("t", r), ("3. C", b), ("t", r)]
    layout = [{"font_size": 11.0, "fonts": f, "bbox": (72, 0, 0, 0)} for _, f in rows]
    heads = [NumberedHeading(str(i), c, 2 * i) for i, c in ((1, "A"), (2, "B"), (3, "C"))]
    assert [h.title for h in _refine_by_bold_headings(heads, [t for t, _ in rows], layout)] == ["A", "B", "C"]
