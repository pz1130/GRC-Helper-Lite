"""边界轮 4：英国 SI 附表、说明注释、GDPR/VDP 等边界修复。"""
from app.parsing.contract import ClauseNode
from app.parsing.pdf_parser import _schedules
from app.parsing.pdf_tables import attach_tables
from app.parsing.restructure import _looks_like_title, nest_chapters


def _node(number, heading, text="", children=None, keys=None, level=1):
    return ClauseNode(heading=heading, text=text, level=level, number=number, kind="section",
                      children=children or [], source_keys=keys or [])


def test_etc_title_is_still_a_title():
    assert _looks_like_title("Criminal offences, penalties and proceedings etc.")
    assert not _looks_like_title("This is a sentence.")


def test_synthetic_part_containers_sort_by_first_child():
    part = _node("Part 1", "Intro", children=[_node("1", "a", keys=["line:10"])])
    sched = _node(None, "SCHEDULES", keys=["line:500"])
    nodes = [sched, part]
    attach_tables(nodes, [])
    assert [n.heading for n in nodes] == ["Intro", "SCHEDULES"]


def test_explanatory_note_after_schedules_is_top_level():
    lines = ["1. Body", "SCHEDULES", "SCHEDULE 1", "Regulation 2", "Bodies", "text",
             "EXPLANATORY NOTE", "(This note is not part)"]
    out = _schedules(lines, 0, [])
    assert out[-1].title == "EXPLANATORY NOTE" and out[-1].line_index == 6


def test_parts_inside_schedule_are_nested():
    sched = _node("SCHEDULE 7", "Amendments", text="PART 1\nPrimary", level=2,
                  children=[_node("1", "Act A", text="x\nPART 2\nSecondary", level=3),
                            _node("2", "Order B", level=3)])
    body = [_node("1", "r1", text="PART 2\nMore"), _node("2", "r2")]
    top = [_node(None, "Unstructured", text="PART 1\nIntro"), *body,
           _node(None, "SCHEDULES", children=[sched])]
    out = nest_chapters(top)
    sch = out[-1].children[0]
    assert [c.number for c in sch.children] == ["Part 1", "Part 2"]
    assert [c.heading for c in sch.children[0].children] == ["Act A"]
    assert [c.heading for c in sch.children[1].children] == ["Order B"]




def test_wrapped_long_title_is_joined_not_demoted():
    from app.parsing.headings import demote_fragment_headings
    node = _node("12", "Transparent information, communication and modalities for the exercise",
                 text="of the rights of the data subject\n1. The controller shall take measures")
    demote_fragment_headings([node])
    assert node.heading.endswith("for the exercise of the rights of the data subject")
    assert node.text.startswith("1. The controller")


def test_sentence_fragment_still_demoted():
    from app.parsing.headings import demote_fragment_headings
    node = _node("2.2", "The company shall ensure", text="that all staff are trained in the handling of data.")
    demote_fragment_headings([node])
    assert node.heading == "2.2"


def test_docx_bullet_items_keep_their_bullet(tmp_path):
    import docx

    from app.parsing.docx_parser import DocxParser
    d = docx.Document()
    d.add_heading("Guidelines", 1)
    d.add_paragraph("Under this policy you:")
    d.add_paragraph("Notify us quickly.", style="List Bullet")
    d.add_paragraph("Avoid harm.", style="List Bullet")
    path = tmp_path / "b.docx"
    d.save(path)
    node = next(c for c in DocxParser().parse(path).clauses if c.heading == "Guidelines")
    assert "• Notify us quickly." in node.text and "• Avoid harm." in node.text
