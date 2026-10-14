from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument
from sqlalchemy import func, select

from app.clauses.models import Clause
from app.ingest.models import DocStatus, DocType, Document
from app.ingest.tasks import run_parse
from app.parsing.contract import ParseError


def _make_docx(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_paragraph("Version: \t1.0")
    doc.add_paragraph("Owner: \tIT Department")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Purpose text.")
    doc.add_heading("Roles and Responsibilities", level=1)
    doc.add_paragraph("IT Division drafts it.")
    path = tmp_path / "g.docx"
    doc.save(path)
    return path


async def _doc(db_session, path: Path, **kwargs) -> Document:
    defaults = {
        "title": "Guideline",
        "doc_type": DocType.GUIDELINE,
        "file_hash": "f" * 64,
        "file_path": str(path),
        "original_filename": path.name,
    }
    defaults.update(kwargs)
    document = Document(**defaults)
    db_session.add(document)
    await db_session.flush()
    return document


@pytest.mark.asyncio
async def test_successful_parse_activates_document(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    result = await run_parse(db_session, document)

    assert document.status is DocStatus.ACTIVE
    assert result["clauses"] >= 2
    assert document.parse_error is None


@pytest.mark.asyncio
async def test_worker_moves_root_tmp_upload_into_document_store(db_session, tmp_path):
    source = tmp_path / "staged.docx"
    _make_docx(tmp_path).replace(source)
    document = await _doc(db_session, source)
    document.file_path = f"/tmp/{source.name}"
    source.rename(document.file_path)

    from app.parsing.contract import DocumentMeta, ParsedDocument

    stored_path = tmp_path / "stored.docx"
    parsed = ParsedDocument(
        meta=DocumentMeta(
            title=None,
            version=None,
            owner=None,
            approver=None,
            approved_date=None,
            effective_date=None,
            doc_type=None,
        ),
        clauses=[],
    )
    with patch("app.ingest.tasks.save") as save_mock, patch(
        "app.ingest.tasks.get_parser"
    ) as get_parser_mock:
        save_mock.return_value.path = str(stored_path)
        get_parser_mock.return_value.parse.return_value = parsed
        await run_parse(db_session, document)

    save_mock.assert_called_once()
    assert document.file_path == str(stored_path)


@pytest.mark.asyncio
async def test_parse_writes_clauses(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    await run_parse(db_session, document)

    headings = [clause.heading for clause in await db_session.scalars(select(Clause))]
    assert "Introduction" in headings
    assert "Roles and Responsibilities" in headings


@pytest.mark.asyncio
async def test_cover_metadata_is_written_back(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    await run_parse(db_session, document)

    assert document.version == "1.0"
    assert document.owner == "IT Department"


@pytest.mark.asyncio
async def test_parse_failure_records_a_specific_reason(db_session, tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")
    document = await _doc(db_session, broken)

    with pytest.raises(Exception):  # noqa: B017 —— 只断言坏文件会失败，具体异常类型不在本测试范围
        await run_parse(db_session, document)

    assert document.status is DocStatus.PARSE_FAILED
    assert document.parse_error
    assert "docx" in document.parse_error.lower()
    # 失败原因原来只有一句中文，界面原样显示；带上 code，前端按 code 查英文词条。
    assert document.parse_error_code == "parse_error.docx_unreadable"
    assert document.parse_error_params == {"error": document.parse_error_params["error"]}


@pytest.mark.asyncio
async def test_an_unsupported_file_type_fails_with_a_code(db_session, tmp_path):
    odd = tmp_path / "policy.txt"
    odd.write_text("1. Scope")
    document = await _doc(db_session, odd)

    with pytest.raises(ParseError):
        await run_parse(db_session, document)

    assert (document.parse_error_code, document.parse_error_params) == (
        "parse_error.unsupported_file_type", {"extension": ".txt"},
    )


@pytest.mark.asyncio
async def test_an_unexpected_failure_still_carries_a_code(db_session, tmp_path, monkeypatch):
    document = await _doc(db_session, _make_docx(tmp_path))

    def explode(path):
        raise RuntimeError("boom")

    monkeypatch.setattr("app.ingest.tasks.get_parser", lambda path: type("P", (), {"parse": staticmethod(explode)})())
    with pytest.raises(RuntimeError):
        await run_parse(db_session, document)

    assert (document.parse_error_code, document.parse_error_params) == (
        "parse_error.unexpected", {"error": "RuntimeError"},
    )


@pytest.mark.asyncio
async def test_completeness_warning_is_stored(db_session, tmp_path, monkeypatch):
    # 期望的章节现在是**部署方配置**，默认什么都不期望——写死那两节是样本语料的
    # 文风，会对每一份外部文档误报（见 test_parsing_validate_house_style.py）。
    monkeypatch.setattr(
        "app.parsing.validate.EXPECTED_SECTIONS", ("Introduction", "Roles and Responsibilities")
    )
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("x")
    path = tmp_path / "partial.docx"
    doc.save(path)

    document = await _doc(db_session, path)
    result = await run_parse(db_session, document)

    assert any("Roles and Responsibilities" in warning for warning in result["warnings"])
    assert "Roles and Responsibilities" in (document.parse_warnings or "")


@pytest.mark.asyncio
async def test_reparse_does_not_duplicate_clauses(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    first = await run_parse(db_session, document)
    second = await run_parse(db_session, document)
    assert first["clauses"] == second["clauses"]

    rows = list(await db_session.scalars(select(Clause)))
    assert len(rows) == first["clauses"]


@pytest.mark.asyncio
async def test_low_confidence_ocr_sets_the_quality_flag(db_session, tmp_path):
    from app.parsing.contract import DocumentMeta, ParsedDocument

    document = await _doc(db_session, _make_docx(tmp_path))
    parsed = ParsedDocument(
        meta=DocumentMeta(
            title=None,
            version=None,
            owner=None,
            approver=None,
            approved_date=None,
            effective_date=None,
            doc_type=None,
        ),
        clauses=[],
        warnings=["OCR 置信度偏低"],
        ocr_used=True,
    )
    with patch("app.ingest.tasks.get_parser") as get_parser_mock:
        get_parser_mock.return_value.parse.return_value = parsed
        await run_parse(db_session, document)

    assert document.ocr_quality_flag is True


@pytest.mark.asyncio
async def test_overlong_cover_field_is_dropped_with_a_warning(db_session, tmp_path):
    # 封面字段是从任意文档的 "Version: ..." 行里切出来的，长度没有上限，列有。
    # 超长直接赋值会在 flush 时让整个解析事务失败——一个元数据字段不该拖垮条款树。
    doc = DocxDocument()
    doc.add_paragraph("Version: \t1.0 (draft for consultation, replaces the 2021 edition)")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Purpose text.")
    path = tmp_path / "long.docx"
    doc.save(path)
    document = await _doc(db_session, path, version="0.9")

    result = await run_parse(db_session, document)

    assert document.status is DocStatus.ACTIVE
    assert document.version == "0.9"
    assert result["clauses"] >= 1
    assert {"code": "parse.cover_field_too_long", "params": {"field": "version", "limit": 32}} in (
        document.parse_warning_codes
    )


@pytest.mark.asyncio
async def test_a_database_failure_after_the_move_leaves_a_reparsable_document(tmp_path):
    """解析之后写库失败：文档要落成 parse_failed，file_path 要指着一个还在的文件。

    曾经：暂存文件在解析**之前**就删了，之后 persist 出 DB 错误 → 事务回滚 →
    file_path 回到那个已删除的 /tmp 路径，状态回到 uploaded。文档永远卡住，
    重新解析只会报 FileNotFoundError，同一文件也传不进来（哈希已被占用）。
    """
    import os
    import uuid
    from unittest.mock import AsyncMock

    from sqlalchemy import text
    from sqlalchemy.exc import DBAPIError
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.ingest.tasks import parse_document

    url = os.environ.get("TEST_DATABASE_URL", "postgresql+asyncpg://grc:grc@db:5432/grc_test")
    engine = create_async_engine(url)
    factory = async_sessionmaker(engine, expire_on_commit=False)

    staged = Path("/tmp") / f"staged-{uuid.uuid4().hex}.docx"
    _make_docx(tmp_path).replace(staged)

    async with factory() as setup:
        document = Document(
            title="stranded probe", doc_type=DocType.GUIDELINE,
            file_hash=uuid.uuid4().hex * 2, file_path=str(staged),
            original_filename="g.docx",
        )
        setup.add(document)
        await setup.commit()
        document_id = document.id

    async def broken_persist(session, nodes, *, document_id, **kwargs):
        await session.execute(text("select 1/0"))

    try:
        with (
            patch("app.ingest.tasks.session_factory", factory),
            patch("app.ingest.tasks.persist", broken_persist),
            patch("app.ingest.storage.DOCUMENT_ROOT", tmp_path / "store"),
            patch("app.worker.enqueue", new=AsyncMock()),
            pytest.raises(DBAPIError),
        ):
            await parse_document({}, document_id)

        async with factory() as check:
            fresh = await check.get(Document, document_id)
            assert fresh.status is DocStatus.PARSE_FAILED
            assert fresh.parse_error
            assert Path(fresh.file_path).is_file()

        # 原因排除之后再解析一次：这回成功，原文进了正式卷，暂存文件被收走。
        with (
            patch("app.ingest.tasks.session_factory", factory),
            patch("app.ingest.storage.DOCUMENT_ROOT", tmp_path / "store"),
            patch("app.worker.enqueue", new=AsyncMock()),
        ):
            await parse_document({}, document_id)

        async with factory() as check:
            fresh = await check.get(Document, document_id)
            assert fresh.status is DocStatus.ACTIVE
            assert Path(fresh.file_path).parent.parent == tmp_path / "store"
            assert Path(fresh.file_path).is_file()
            assert not staged.exists()
    finally:
        staged.unlink(missing_ok=True)
        async with factory() as cleanup:
            stale = await cleanup.get(Document, document_id)
            if stale is not None:
                await cleanup.delete(stale)
                await cleanup.commit()
        await engine.dispose()


async def _previous_version(db_session, tmp_path) -> Document:
    from app.controls.models import Control, ControlSource, SourceRelation

    previous = await _doc(
        db_session, tmp_path / "old.docx", file_hash="0" * 64, status=DocStatus.ACTIVE
    )
    clause = Clause(
        document_id=previous.id, number="1", heading="Scope", heading_path="Scope",
        citation_label="1", text="Access shall be reviewed.", order_index=0, level=1,
    )
    control = Control(code="AC-OLD", title="Access review")
    db_session.add_all([clause, control])
    await db_session.flush()
    db_session.add(
        ControlSource(control_id=control.id, clause_id=clause.id, relation=SourceRelation.DEFINES)
    )
    await db_session.flush()
    return previous


@pytest.mark.asyncio
async def test_a_parsed_new_version_supersedes_the_previous_one(db_session, tmp_path):
    from app.controls.models import ControlSource
    from app.iam.models import AuditLog

    previous = await _previous_version(db_session, tmp_path)
    document = await _doc(db_session, _make_docx(tmp_path), supersedes_id=previous.id)

    await run_parse(db_session, document)

    assert previous.status is DocStatus.SUPERSEDED
    # 旧版上已确认的出处原样留着：那是历史，不是垃圾。
    assert await db_session.scalar(select(func.count()).select_from(ControlSource)) == 1
    entry = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "document.superseded")
    )
    assert entry.entity_id == str(previous.id)
    assert entry.before == {"status": "active"}
    assert entry.after == {"status": "superseded", "superseded_by": document.id}


@pytest.mark.asyncio
async def test_a_failed_new_version_leaves_the_previous_one_active(db_session, tmp_path):
    # 新版解析失败时旧版仍是唯一可用的现行版本，不能先把它撤下来。
    previous = await _previous_version(db_session, tmp_path)
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")
    document = await _doc(db_session, broken, supersedes_id=previous.id)

    with pytest.raises(Exception):  # noqa: B017 - 这里只关心旧版的状态
        await run_parse(db_session, document)

    assert previous.status is DocStatus.ACTIVE
